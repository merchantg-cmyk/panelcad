"""PLC I/O list report and Studio 5000 / TIA / Do-more tag export."""

from __future__ import annotations

import csv
from pathlib import Path

import pytest

from panelcad import Session
from panelcad.cli.agent import Agent
from panelcad.core.plc_tags import (
    CXPROGRAMMER_DOC,
    DOMORE_DOC,
    KEYENCE_SPECS,
    STUDIO5000_DOC,
    SYSMAC_DOC,
    TIA_PLC_COLUMNS,
    TIA_PROP_COLUMNS,
    TIA_XLSX_DOC,
    keyence_address_collisions,
    omron_address_collisions,
    render_tags,
    tia_byte_collisions,
)
from panelcad.render.print_package import extract_pdf_text, plan_print_package

GOLDEN = Path(__file__).parent / "golden" / "plc_io"


def _assert_tia_workbook(path: Path, *, connected_name: str, connected_address: str) -> None:
    from openpyxl import load_workbook

    book = load_workbook(path)
    assert book.sheetnames == ["PLCTags", "TagTableProperties"]
    tags = book["PLCTags"]
    assert [cell.value for cell in tags[1]] == list(TIA_PLC_COLUMNS)
    body = [row for row in tags.iter_rows(min_row=2, values_only=True) if row[0]]
    hit = next(row for row in body if row[0] == connected_name)
    assert hit[1] == "PLC tags"
    assert hit[3] == connected_address
    assert hit[5] == "TRUE"
    assert hit[6] == "TRUE"
    assert hit[7] in {"TRUE", "FALSE"}
    props = book["TagTableProperties"]
    assert [cell.value for cell in props[1]] == list(TIA_PROP_COLUMNS)
    assert props["A2"].value == "PLC tags"
    custom = {item.name: item.value for item in book.custom_doc_props}
    assert custom["TIA_Version"] == "1.0"


def _job() -> Session:
    session = Session.create("PLC IO")
    session.place("iec.plc.compact_16dio", ref="PLC1", x=80, y=40)
    session.place("iec.operators.pb_no", ref="S1", x=20, y=40)
    session.place("iec.contactors.coil", ref="CR1", x=160, y=40)
    session.set_component_properties("S1", {"function": "Start"})
    session.set_component_properties("CR1", {"function": "Run"})
    session.wire("S1.14", "PLC1.I0_0")
    session.wire("PLC1.Q0_0", "CR1.A1")
    session.set_wire_spec(frm="S1.14", to="PLC1.I0_0", number="101")
    session.set_wire_spec(frm="PLC1.Q0_0", to="CR1.A1", number="102")
    return session


def _qt_app():
    from PySide6.QtWidgets import QApplication

    if QApplication.instance() is None:
        QApplication([])


def test_io_rows_describe_connected_points_and_spares():
    rows = {row["address"]: row for row in _job().io_rows()}
    assert "Lplus" not in rows
    assert "M" not in rows
    start = rows["I0_0"]
    assert start["point"] == "I0.0"
    assert start["type"] == "DI"
    assert start["tag"] == "S1"
    assert start["device"] == "S1"
    assert start["function"] == "Start"
    assert start["wire"] == "101"
    assert start["spare"] == "no"
    assert start["description"] == "S1, Start, wire 101, sheet main line 104"
    assert rows["Q0_0"]["type"] == "DO"
    assert rows["Q0_0"]["tag"] == "CR1"
    assert rows["Q0_0"]["description"] == "CR1, Run, wire 102, sheet main line 103"
    spare = rows["I0_1"]
    assert spare["spare"] == "yes"
    assert spare["tag"] == "Spare_I0_1"
    assert spare["device"] == ""
    assert spare["description"] == "spare"
    assert spare["point"] == "I0.1"


def test_io_module_is_listed_with_the_plc():
    session = Session.create("Modules")
    session.place("iec.plc.compact_16dio", ref="PLC1", x=40, y=40)
    session.place("iec.plc.di16", ref="A2", x=140, y=40)
    session.set_component_properties("A2", {"slot": "5"})
    rows = session.io_rows()
    modules = {row["module"] for row in rows}
    assert modules == {"PLC1", "A2"}
    extra = [row for row in rows if row["module"] == "A2" and row["address"] == "I0"]
    assert extra[0]["type"] == "DI"
    assert extra[0]["slot"] == "5"
    assert extra[0]["spare"] == "yes"


def test_exporter_goldens_match_byte_for_byte():
    rows = _job().io_rows()
    studio = render_tags(rows, target="studio5000")
    tia = render_tags(rows, target="tia_csv")
    domore = render_tags(rows, target="domore")
    assert studio == (GOLDEN / "studio5000-tags.csv").read_text(encoding="utf-8")
    assert tia == (GOLDEN / "tia-tags.csv").read_text(encoding="utf-8")
    assert domore == (GOLDEN / "domore-tags.csv").read_text(encoding="utf-8")
    assert studio.startswith("0.3\nRemark,")
    assert "Not a TIA Portal import file" in tia
    assert tia.splitlines()[5] == "name,data_type,logical_address,comment"
    parsed = list(csv.reader(domore.splitlines()))
    assert parsed[0] == ["X0", "S1", "101", "S1 Start wire 101 sheet main line 104"]
    assert parsed[4] == ["Y0", "CR1", "102", "CR1 Run wire 102 sheet main line 103"]
    assert all(len(row) == 4 for row in parsed)


def test_studio5000_records_have_seven_fields_and_blank_alias_type():
    text = (GOLDEN / "studio5000-tags.csv").read_text(encoding="utf-8")
    rows = list(csv.reader(text.splitlines()))
    assert rows[0] == ["0.3"]
    assert rows[1][0] == "Remark"
    assert rows[2] == ["TYPE", "SCOPE", "NAME", "DESCRIPTION", "DATATYPE", "SPECIFIER", "ATTRIBUTES"]
    data = rows[3]
    assert data[0] == "ALIAS"
    assert data[1] == ""
    assert data[2] == "S1"
    assert data[4] == ""
    assert data[5] == "Local:1:I.Pt00.Data"
    assert data[6] == ""
    assert len(data) == 7
    outputs = [row for row in rows if row[0] == "ALIAS" and row[2] == "CR1"]
    assert outputs[0][5] == "Local:1:O.Pt00.Data"


def test_naming_prefix_description_and_base_tags():
    rows = _job().io_rows()
    base = render_tags(rows, target="studio5000", prefix="M1", name_from="description", logix_style="base")
    assert "TAG,,M1_S1_Start_wire_101_sheet_main_line_104," in base
    assert ",BOOL,," in base
    assert "ALIAS" not in base
    nick = render_tags(rows, target="domore", prefix="M1", name_from="device").splitlines()[0]
    assert nick.startswith('"X0","M1_S1",')
    long_name = next(csv.reader(render_tags(rows, target="domore", name_from="description").splitlines()))
    assert long_name[0] == "X0"
    assert 1 <= len(long_name[1]) <= 16
    assert "," not in long_name[3]


def test_second_module_and_analog_addresses():
    session = Session.create("Two")
    session.place("iec.plc.compact_16dio", ref="PLC1", x=40, y=40)
    session.place("iec.plc.di16", ref="A2", x=140, y=40)
    session.set_component_properties("A2", {"slot": "5"})
    rows = session.io_rows()
    studio = render_tags(rows, target="studio5000")
    tia = render_tags(rows, target="tia_csv")
    domore = render_tags(rows, target="domore")
    assert "Local:1:I.Pt00.Data" in studio
    assert "Local:5:I.Pt00.Data" in studio
    assert "%I0.0" in tia
    assert "%I1.0" in tia
    assert '"X0"' in domore
    assert '"X8"' in domore
    analog = [
        {
            "module": "AI1",
            "address": "AI0",
            "type": "AI",
            "description": "Level",
            "wire": "201",
            "device": "LT1",
            "tag": "LT1",
            "spare": "no",
            "slot": "",
        },
        {
            "module": "AO1",
            "address": "AO1",
            "type": "AO",
            "description": "Valve",
            "wire": "202",
            "device": "CV1",
            "tag": "CV1",
            "spare": "no",
            "slot": "",
        },
        {
            "module": "DI1",
            "address": "X0",
            "type": "DI",
            "description": "Start",
            "wire": "1",
            "device": "S1",
            "tag": "S1",
            "spare": "no",
            "slot": "3",
        },
    ]
    studio_a = render_tags(analog, target="studio5000")
    tia_a = render_tags(analog, target="tia_csv")
    domore_a = render_tags(analog, target="domore")
    assert "Local:1:I.Ch00.Data" in studio_a
    assert "Local:2:O.Ch01.Data" in studio_a
    assert "Local:3:I.Pt00.Data" in studio_a
    assert "REAL" not in studio_a
    assert "%IW2" in tia_a
    assert "%IW0" not in tia_a
    assert "%QW2" in tia_a
    assert '"WX0","LT1"' in domore_a
    assert '"WY1","CV1"' in domore_a
    assert '"X0","S1"' in domore_a
    base = render_tags(analog, target="logix", logix_style="base")
    assert ",REAL,," in base


def test_nema_plc_motor_sample_includes_analog_channels(tmp_path: Path):
    from panelcad.examples.plc_io import build_plc_motor_io

    session = build_plc_motor_io()
    rows = session.io_rows()
    kinds = {row["type"] for row in rows}
    assert {"DI", "DO", "AI", "AO"} <= kinds
    assert sum(row["type"] == "DI" for row in rows) >= 4
    by_address = {row["address"]: row for row in rows}
    assert by_address["I0_0"]["device"] == "M1"
    assert by_address["I0_0"]["tag"] == "M1_RUN_FB"
    assert by_address["I0_1"]["device"] == "OL1"
    assert by_address["I0_1"]["tag"] == "OL1_OK"
    assert by_address["I0_2"]["device"] == "SEL1"
    assert by_address["I0_2"]["tag"] == "SEL1_AUTO"
    assert by_address["AI0"]["tag"] == "LT1" and by_address["AI0"]["device"] == "LT1"
    assert by_address["AO0"]["tag"] == "CV1" and by_address["AO0"]["device"] == "CV1"
    assert by_address["I0_1"]["line"] == "309"
    assert by_address["AI0"]["line"] == "314"
    assert by_address["AO0"]["line"] == "316"
    assert len({by_address["I0_1"]["line"], by_address["AI0"]["line"], by_address["AO0"]["line"]}) == 3
    studio = render_tags(rows, target="studio5000")
    lines = studio.splitlines()
    assert lines[0] == "0.3"
    assert lines[1].startswith('Remark,"')
    assert lines[2].startswith("TYPE,")
    assert "Local:1:I.Pt00.Data" in studio
    assert "Local:2:I.Ch00.Data" in studio
    assert "Local:2:O.Ch00.Data" in studio
    domore = render_tags(rows, target="domore")
    assert '"WX0","LT1"' in domore
    assert '"WY0","CV1"' in domore
    book = tmp_path / "motor.xlsx"
    result = session.export_plc_tags(book, target="tia")
    assert result["count"] == len(rows)
    _assert_tia_workbook(book, connected_name="LT1", connected_address="%IW2")


def _tia_addresses(text: str) -> list[str]:
    body = [line for line in text.splitlines() if line and not line.startswith("#")]
    return [row[2] for row in csv.reader(body[1:])]


def _domore_elements(text: str) -> list[str]:
    return [row[0] for row in csv.reader(text.splitlines()) if row]


def _logix_specifiers(text: str) -> list[str]:
    rows = list(csv.reader(text.splitlines()))
    return [row[5] for row in rows if row and row[0] in {"ALIAS", "TAG"} and len(row) > 5 and row[5]]


def test_tia_analog_words_do_not_share_digital_bytes(tmp_path: Path):
    import re

    from panelcad.examples.plc_io import build_plc_motor_io

    session = build_plc_motor_io()
    rows = session.io_rows()
    tia = render_tags(rows, target="tia_csv")
    addresses = _tia_addresses(tia)
    assert "%IW2" in addresses
    assert "%QW2" in addresses
    assert "%IW0" not in addresses
    assert tia_byte_collisions(addresses) == []
    assert tia_byte_collisions(["%I0.0", "%I0.7"]) == []
    assert tia_byte_collisions(["%I0.0", "%IW0"])
    assert tia_byte_collisions(["%Q0.3", "%QW0"])

    mixed = [
        {"module": "DI1", "address": "I0_0", "type": "DI", "tag": "S1", "description": "Start", "spare": "no"},
        {"module": "DI1", "address": "I1_0", "type": "DI", "tag": "S2", "description": "Stop", "spare": "no"},
        {"module": "AI1", "address": "AI0", "type": "AI", "tag": "LT1", "description": "Level", "spare": "no"},
        {"module": "AO1", "address": "AO1", "type": "AO", "tag": "CV1", "description": "Speed", "spare": "no"},
    ]
    mixed_tia = _tia_addresses(render_tags(mixed, target="tia_csv"))
    assert "%I0.0" in mixed_tia and "%I1.0" in mixed_tia
    assert "%IW2" in mixed_tia
    assert "%QW2" in mixed_tia
    assert tia_byte_collisions(mixed_tia) == []

    forced = _tia_addresses(render_tags(mixed, target="tia_csv", analog_start=80))
    assert "%IW80" in forced and "%QW82" in forced
    odd = _tia_addresses(render_tags(mixed, target="tia_csv", analog_start=81))
    assert "%IW82" in odd and "%QW84" in odd

    owned = [dict(row) for row in mixed]
    for row in owned:
        if row["type"] in {"AI", "AO"}:
            row["start_address"] = "64"
    library = _tia_addresses(render_tags(owned, target="tia_csv", analog_start=80))
    assert "%IW64" in library and "%QW66" in library
    assert "%IW80" not in library and "%QW82" not in library

    device = session.catalog.get("nema.plc.analog_1ai_1ao")
    saved = dict(device.properties)
    device.properties["start_address"] = 128
    try:
        tagged = _tia_addresses(render_tags(session.io_rows(), target="tia_csv", analog_start=80))
    finally:
        device.properties.clear()
        device.properties.update(saved)
    assert "%IW128" in tagged and "%QW128" in tagged

    session.project.properties["analog_start"] = 96
    project_text = session.export_plc_tags(tmp_path / "analog-start.csv", target="tia_csv")
    exported = _tia_addresses(Path(project_text["path"]).read_text(encoding="utf-8"))
    assert "%IW96" in exported and "%QW96" in exported

    domore = _domore_elements(render_tags(rows, target="domore"))
    assert "X0" in domore and "WX0" in domore
    assert len(domore) == len(set(domore))
    kinds: dict[str, set[int]] = {}
    for element in domore:
        match = re.fullmatch(r"(WX|WY|X|Y)(\d+)", element)
        assert match
        kinds.setdefault(match.group(1), set()).add(int(match.group(2)))
    assert all(len(values) == len(set(values)) for values in kinds.values())

    specifiers = _logix_specifiers(render_tags(rows, target="studio5000"))
    assert len(specifiers) == len(set(specifiers))
    keys = []
    for spec in specifiers:
        match = re.fullmatch(r"Local:(\d+):([IO])\.(?:Pt(\d+)\.Data|Ch(\d+)\.Data)", spec)
        assert match, spec
        if match.group(3) is not None:
            keys.append((match.group(1), match.group(2), "pt", int(match.group(3))))
        else:
            keys.append((match.group(1), match.group(2), "ch", int(match.group(4))))
    assert len(keys) == len(set(keys))
    assert ("2", "I", "ch", 0) in keys
    assert ("1", "I", "pt", 0) in keys


def test_contact_tags_follow_parent_role_and_vendor_rules():
    import re

    from panelcad.core.plc_tags import _ELEMENT_ID

    session = Session.create("Contacts")
    session.place("nema.plc.compact_16dio", ref="PLC1", x=160, y=40)
    session.place("nema.contactors.aux_no", ref="OL1-13", x=20, y=40)
    session.place("nema.contactors.aux_no", ref="M1-14", x=20, y=70)
    session.place("nema.contactors.aux_no", ref="M1-13", x=20, y=100)
    session.place("nema.operators.selector_3pos", ref="SEL1", x=20, y=130)
    session.set_component_properties("OL1-13", {"parent": "OL1", "role": "ol_sense"})
    session.set_component_properties("M1-14", {"parent": "M1", "role": "run_fb"})
    session.set_component_properties("M1-13", {"parent": "M1", "role": "run_fb"})
    session.set_component_properties("SEL1", {"pin_labels": {"2": "H", "3": "A"}})
    session.wire("OL1-13.4", "PLC1.I0_0", number="1")
    session.wire("M1-14.4", "PLC1.I0_1", number="2")
    session.wire("M1-13.4", "PLC1.I0_2", number="3")
    session.wire("SEL1.3", "PLC1.I0_3", number="4")
    rows = {row["address"]: row for row in session.io_rows()}
    assert rows["I0_0"]["device"] == "OL1"
    assert rows["I0_0"]["tag"] == "OL1_TRIP"
    assert rows["I0_1"]["tag"] == "M1_RUN_FB"
    assert rows["I0_2"]["tag"] == "M1_RUN_FB_2"
    assert rows["I0_3"]["device"] == "SEL1"
    assert rows["I0_3"]["tag"] == "SEL1_AUTO"

    described = render_tags(session.io_rows(), target="studio5000", name_from="description")
    described_names = [row[2] for row in csv.reader(described.splitlines()) if row and row[0] == "ALIAS"]
    assert any(name.startswith("OL1_Overload_trip") for name in described_names)
    assert "OL1_TRIP" not in described_names
    assert len({name.casefold() for name in described_names}) == len(described_names)

    limits = {"studio5000": 40, "tia_csv": 40, "domore": 16}
    motor = Session.create("Names")
    from panelcad.examples.plc_io import build_plc_motor_io

    motor_rows = build_plc_motor_io().io_rows()
    for target, limit in limits.items():
        text = render_tags(motor_rows, target=target, name_from="device")
        if target == "domore":
            found = [row[1] for row in csv.reader(text.splitlines())]
        elif target == "studio5000":
            found = [row[2] for row in csv.reader(text.splitlines()) if row and row[0] == "ALIAS"]
        else:
            body = [line for line in text.splitlines() if line and not line.startswith("#")]
            found = [row[0] for row in csv.reader(body[1:])]
        assert len(found) == len({name.casefold() for name in found})
        for name in found:
            assert re.fullmatch(r"[A-Za-z_][A-Za-z0-9_]*", name), name
            assert len(name) <= limit
            if target == "domore":
                assert _ELEMENT_ID.match(name) is None

    long_rows = [
        {
            "module": "PLC1",
            "address": "I0_0",
            "type": "DI",
            "tag": "1START_" + ("A" * 50),
            "description": "Start",
            "spare": "no",
        }
    ]
    studio_name = next(
        row[2] for row in csv.reader(render_tags(long_rows, target="studio5000").splitlines()) if row[0] == "ALIAS"
    )
    assert studio_name.startswith("T_")
    assert len(studio_name) <= 40
    nick = next(csv.reader(render_tags(long_rows, target="domore").splitlines()))[1]
    assert len(nick) <= 16
    assert nick[0].isalpha() or nick[0] == "_"


def test_io_line_matches_the_cross_reference_rung():
    from panelcad.core.xref import component_location
    from panelcad.examples.plc_io import build_plc_motor_io

    session = build_plc_motor_io()
    sheet = session.project.sheet("plc-io")
    rows = {row["address"]: row for row in session.io_rows()}
    for ref, address in (("OL1-96", "I0_1"), ("LT1", "AI0"), ("CV1", "AO0"), ("M1-14", "I0_0")):
        comp = sheet.component(ref)
        device = session.catalog.get(comp.library_id)
        location = component_location(session.project, sheet, comp, device)
        assert rows[address]["line"] == str(location.line)
        assert rows[address]["device_sheet"] == "plc-io"
    assert rows["I0_1"]["line"] != rows["AI0"]["line"]
    assert rows["AI0"]["line"] != rows["AO0"]["line"]


def test_click_native_x_addresses_stay_put():
    session = Session.create("CLICK")
    session.place("adc.click.c0_16nd3", ref="A1", x=40, y=40)
    text = render_tags(session.io_rows(), target="brx")
    first = next(csv.reader(text.splitlines()))
    assert first[0] == "X0"
    studio = render_tags(session.io_rows(), target="rockwell")
    assert "Local:1:I.Pt00.Data" in studio


def test_export_files_and_rejects_unknown_target(tmp_path: Path):
    session = _job()
    dest = tmp_path / "nested" / "tags.xlsx"
    result = session.export_plc_tags(dest, target="siemens")
    assert result["target"] == "tia"
    assert result["verified"] is True
    assert result["format"] == "tia-xlsx"
    assert result["doc"] == TIA_XLSX_DOC
    _assert_tia_workbook(dest, connected_name="S1", connected_address="%I0.0")
    generic = session.export_plc_tags(tmp_path / "generic.csv", target="tia_csv")
    assert generic["verified"] is False
    assert generic["format"] == "tia-generic-csv"
    assert (tmp_path / "generic.csv").read_text(encoding="utf-8") == (GOLDEN / "tia-tags.csv").read_text(
        encoding="utf-8"
    )
    studio = session.export_plc_tags(tmp_path / "s.csv", target="studio5000")
    assert studio["verified"] is True
    assert studio["doc"] == STUDIO5000_DOC
    domore = session.export_plc_tags(tmp_path / "d.csv", target="do_more")
    assert domore["verified"] is True
    assert domore["doc"] == DOMORE_DOC
    with pytest.raises(ValueError, match="unknown PLC tag target"):
        session.export_plc_tags(tmp_path / "bad.csv", target="modbus")


def test_print_package_plc_io_section(tmp_path: Path):
    _qt_app()
    session = _job()
    plan = plan_print_package(session, sections=["plc_io"])
    assert plan.sections == ["plc_io"]
    assert plan.pages[0].title == "PLC I/O List"
    aliased = session.plan_print_package(sections="io_list")
    assert aliased["sections"] == ["plc_io"]
    empty = Session.create("Empty").plan_print_package()
    assert "plc_io" not in empty["sections"]
    dest = tmp_path / "io.pdf"
    session.print_package(dest, sections=["plc_io"])
    text = extract_pdf_text(dest, 0)
    assert "PLC I/O LIST" in text
    assert "I0.0" in text
    assert "S1" in text
    assert "spare" in text


def test_gui_menus_dialog_and_export_actions(tmp_path: Path, monkeypatch: pytest.MonkeyPatch):
    from PySide6.QtWidgets import QApplication, QMessageBox

    from panelcad.gui.main_window import MainWindow
    from panelcad.gui.plc_tag_dialog import PlcTagExportDialog

    _qt_app()
    app = QApplication.instance()
    window = MainWindow(_job())
    window.show()
    app.processEvents()

    def labeled(actions, title: str):
        for action in actions:
            if action.text().replace("&", "") == title:
                return action
        raise AssertionError(title)

    file_menu = labeled(window.menuBar().actions(), "File").menu()
    export_menu = labeled(file_menu.actions(), "Export").menu()
    tags_action = labeled(export_menu.actions(), "PLC Tags…")
    reports = labeled(window.menuBar().actions(), "Reports").menu()
    io_action = labeled(reports.actions(), "PLC I/O List…")

    dialog = PlcTagExportDialog()
    assert dialog.values()["target"] == "studio5000"
    assert dialog.values()["name_from"] == "device"
    assert dialog.values()["analog_start"] is None
    dialog.analog_start.setValue(64)
    assert dialog.values()["analog_start"] == 64
    dialog.analog_start.setValue(81)
    assert dialog.values()["analog_start"] == 82
    dialog.analog_start.setValue(-1)
    assert str(dialog.values()["path"]).endswith("studio5000-tags.csv")
    assert dialog.logix_style.isEnabled()
    dialog.target.setCurrentIndex(1)
    assert dialog.values()["target"] == "tia"
    assert "PLCTags" in dialog.note.text()
    assert dialog.logix_style.isEnabled() is False
    assert str(dialog.values()["path"]).endswith("tia-tags.xlsx")
    dialog.target.setCurrentIndex(2)
    assert dialog.values()["target"] == "tia_csv"
    assert "not a TIA import" in dialog.note.text()
    dialog.path.setText(str(tmp_path / "custom.csv"))
    dialog.target.setCurrentIndex(3)
    assert dialog.values()["target"] == "domore"
    assert str(dialog.values()["path"]).endswith("custom.csv")
    dialog.prefix.setText("M1")
    dialog.name_from.setCurrentIndex(1)
    dialog.logix_slot.setValue(4)
    vals = dialog.values()
    assert vals["prefix"] == "M1"
    assert vals["name_from"] == "description"
    assert vals["logix_slot"] == 4
    dialog.close()

    io_path = tmp_path / "io-list.csv"
    monkeypatch.setattr(
        "panelcad.gui.main_window.QFileDialog.getSaveFileName",
        lambda *args, **kwargs: (str(io_path), "CSV (*.csv)"),
    )
    io_action.trigger()
    app.processEvents()
    io_text = io_path.read_text(encoding="utf-8")
    assert "address" in io_text.splitlines()[0]
    assert "I0_0" in io_text
    assert "spare" in io_text

    tag_path = tmp_path / "studio.csv"
    monkeypatch.setattr(PlcTagExportDialog, "exec", lambda self: self.DialogCode.Accepted)
    monkeypatch.setattr(
        PlcTagExportDialog,
        "values",
        lambda self: {
            "path": str(tag_path),
            "target": "studio5000",
            "prefix": "",
            "name_from": "device",
            "logix_style": "alias",
            "logix_slot": 1,
        },
    )
    tags_action.trigger()
    app.processEvents()
    assert tag_path.read_text(encoding="utf-8") == (GOLDEN / "studio5000-tags.csv").read_text(encoding="utf-8")

    notes: list[str] = []
    monkeypatch.setattr(QMessageBox, "information", lambda *args, **kwargs: notes.append(str(args[-1])))
    monkeypatch.setattr(
        PlcTagExportDialog,
        "values",
        lambda self: {
            "path": str(tmp_path / "tia.csv"),
            "target": "tia_csv",
            "prefix": "",
            "name_from": "device",
            "logix_style": "alias",
            "logix_slot": 1,
        },
    )
    tags_action.trigger()
    assert notes
    assert "not a TIA import" in notes[-1]
    window.close()


def test_jsonl_schema_and_commands(tmp_path: Path):
    agent = Agent()
    help_result = agent.handle({"id": 1, "op": "help"})
    assert help_result["ok"]
    schemas = help_result["result"]["schemas"]
    plc = schemas["plc_io_list"]
    tags = schemas["export_plc_tags"]
    assert "spares" in plc["description"]
    assert plc["args"]["path"]["description"]
    assert tags["args"]["target"]["enum"] == [
        "studio5000",
        "tia",
        "tia_csv",
        "domore",
        "keyence",
        "sysmac",
        "cxprogrammer",
    ]
    for name in ("path", "target", "prefix", "name_from", "logix_style", "logix_slot", "analog_start"):
        assert tags["args"][name]["description"]
        assert name in help_result["result"]["help"]["export_plc_tags"]
    assert "plc_io_list" in help_result["result"]["ops"]
    assert "export_plc_tags" in help_result["result"]["ops"]

    assert agent.handle({"id": 2, "op": "create_project", "name": "Agent PLC"})["ok"]
    assert agent.handle(
        {"id": 3, "op": "place", "library_id": "iec.plc.compact_16dio", "ref": "PLC1", "x": 80, "y": 40}
    )["ok"]
    assert agent.handle(
        {"id": 4, "op": "place", "library_id": "iec.operators.pb_no", "ref": "S1", "x": 20, "y": 40}
    )["ok"]
    assert agent.handle({"id": 5, "op": "set_properties", "ref": "S1", "function": "Start"})["ok"]
    wired = agent.handle({"id": 6, "op": "wire", "from": "S1.14", "to": "PLC1.I0_0", "number": "101"})
    assert wired["ok"], wired
    listed = agent.handle({"id": 7, "op": "plc_io_list", "path": str(tmp_path / "agent-io.csv")})
    assert listed["ok"], listed
    assert listed["result"]["count"] == 8
    start = next(row for row in listed["result"]["rows"] if row["address"] == "I0_0")
    assert start["type"] == "DI"
    assert start["tag"] == "S1"
    assert start["spare"] == "no"
    assert "wire 101" in start["description"]
    assert (tmp_path / "agent-io.csv").is_file()
    exported = agent.handle(
        {
            "id": 8,
            "op": "export_plc_tags",
            "path": str(tmp_path / "agent-tags.csv"),
            "target": "domore",
            "prefix": "M1",
            "name_from": "device",
        }
    )
    assert exported["ok"], exported
    assert exported["result"]["format"] == "domore-element-documentation"
    assert '"X0","M1_S1"' in (tmp_path / "agent-tags.csv").read_text(encoding="utf-8")
    rejected = agent.handle({"id": 9, "op": "export_plc_tags", "path": str(tmp_path / "x.csv"), "target": "abb"})
    assert rejected["ok"] is False
    assert "target" in rejected["error"]


def test_drawing_agent_tool_schema_and_export(tmp_path: Path):
    from pydantic_ai.messages import ModelMessage, ModelResponse, TextPart, ToolCallPart
    from pydantic_ai.models.function import AgentInfo, FunctionModel

    from panelcad.ai.drawing import drawing_agent, open_drawing

    tools: dict[str, object] = {}

    def capture(messages: list[ModelMessage], info: AgentInfo) -> ModelResponse:
        tools.update({tool.name: tool for tool in info.function_tools})
        return ModelResponse(parts=[TextPart(content="ready")])

    drawing = open_drawing(name="PLC tools")
    with drawing_agent.override(model=FunctionModel(capture)):
        result = drawing_agent.run_sync("List PLC tools.", deps=drawing)
    assert result.output == "ready"
    io_tool = tools["plc_io_list"]
    tag_tool = tools["export_plc_tags"]
    assert "spare" in (io_tool.description or "").lower()
    tag_desc = (tag_tool.description or "").lower()
    assert "studio 5000" in tag_desc
    assert "do-more" in tag_desc
    props = tag_tool.parameters_json_schema["properties"]
    assert "generic" in props["target"]["description"].lower()
    assert set(props["target"]["enum"]) == {
        "studio5000",
        "tia",
        "tia_csv",
        "domore",
        "keyence",
        "sysmac",
        "cxprogrammer",
    }
    assert props["prefix"]["description"]
    assert set(props["name_from"]["enum"]) == {"device", "description"}
    assert "analog" in props["analog_start"]["description"].lower()
    assert set(props["logix_style"]["enum"]) == {"alias", "base"}
    required = set(tag_tool.parameters_json_schema.get("required") or [])
    assert {"path", "target"} <= required

    session = _job()
    src = tmp_path / "job.pcad.yaml"
    session.save(src)
    drawing = open_drawing(src)
    dest = tmp_path / "from-agent.csv"
    steps = [
        ModelResponse(
            parts=[
                ToolCallPart(
                    tool_name="export_plc_tags",
                    args={"path": str(dest), "target": "studio5000"},
                    tool_call_id="export-tags",
                )
            ]
        ),
        ModelResponse(parts=[TextPart(content="Exported Studio 5000 tags.")]),
    ]

    def caller(messages: list[ModelMessage], info: AgentInfo) -> ModelResponse:
        return steps.pop(0)

    with drawing_agent.override(model=FunctionModel(caller)):
        exported = drawing_agent.run_sync("Export the PLC tags.", deps=drawing)
    assert "Studio 5000" in exported.output
    assert dest.read_text(encoding="utf-8") == (GOLDEN / "studio5000-tags.csv").read_text(encoding="utf-8")


def test_cli_plc_io_and_plc_tags(tmp_path: Path):
    from typer.testing import CliRunner

    from panelcad.cli.main import app

    src = tmp_path / "job.pcad.yaml"
    _job().save(src)
    runner = CliRunner()
    listed = runner.invoke(app, ["electrical", "plc-io", str(src)])
    assert listed.exit_code == 0, listed.output
    assert "I0.0" in listed.output
    assert "S1" in listed.output
    assert "spare" in listed.output
    out = tmp_path / "cli-io.csv"
    saved = runner.invoke(app, ["electrical", "plc-io", str(src), "--out", str(out)])
    assert saved.exit_code == 0, saved.output
    assert "I0_0" in out.read_text(encoding="utf-8")
    tags = tmp_path / "cli-tags.csv"
    exported = runner.invoke(
        app,
        ["electrical", "plc-tags", str(src), "--out", str(tags), "--target", "studio5000"],
    )
    assert exported.exit_code == 0, exported.output
    assert "studio5000-csv" in exported.output
    assert tags.read_text(encoding="utf-8") == (GOLDEN / "studio5000-tags.csv").read_text(encoding="utf-8")
    bad = runner.invoke(app, ["electrical", "plc-tags", str(src), "--out", str(tmp_path / "bad.csv"), "--target", "abb"])
    assert bad.exit_code == 1
    assert "unknown PLC tag target" in bad.output
    cx_out = tmp_path / "cx.csv"
    cx_cli = runner.invoke(
        app,
        ["electrical", "plc-tags", str(src), "--out", str(cx_out), "--target", "cxprogrammer"],
    )
    assert cx_cli.exit_code == 0, cx_cli.output
    assert cx_cli.output.strip().endswith(".txt") or "cxprogrammer-symbol-txt" in cx_cli.output
    assert (tmp_path / "cx.txt").is_file()


def _data_rows(text: str) -> list[list[str]]:
    return list(csv.reader(line for line in text.splitlines() if line and not line.startswith("#")))


def test_keyence_omron_addresses_do_not_overlap(tmp_path: Path):
    from panelcad.examples.plc_io import build_plc_motor_io

    rows = build_plc_motor_io().io_rows()
    cx = render_tags(rows, target="cxprogrammer")
    parsed = [line.split("\t") for line in cx.splitlines() if line]
    assert all(len(row) == 4 for row in parsed)
    assert not cx.startswith("#")
    by_name = {row[0]: row for row in parsed}
    assert by_name["M1_RUN_FB"] == [
        "M1_RUN_FB",
        "BOOL",
        "0.00",
        by_name["M1_RUN_FB"][3],
    ]
    assert by_name["OL1_OK"][2] == "0.01"
    assert by_name["SEL1_AUTO"][2] == "0.02"
    assert by_name["CR1"][1:] == ["BOOL", "100.00", by_name["CR1"][3]]
    assert by_name["LT1"][1] == "UINT"
    assert by_name["LT1"][2] == "2005"
    assert by_name["CV1"][2] == "2001"
    cx_addresses = [row[2] for row in parsed]
    assert omron_address_collisions(cx_addresses) == []
    assert omron_address_collisions(["0.00", "0.01", "2001"]) == []
    assert omron_address_collisions(["2001.00", "2001"])
    assert omron_address_collisions(["0.00", "0.00"])

    keyence = render_tags(rows, target="kv")
    assert "Not a KV STUDIO import file" in keyence
    krows = _data_rows(keyence)
    assert krows[0] == ["name", "device", "type", "comment"]
    kmap = {row[0]: row[1] for row in krows[1:]}
    assert kmap["M1_RUN_FB"] == "R30000"
    assert kmap["OL1_OK"] == "R30001"
    assert kmap["SEL1_AUTO"] == "R30002"
    assert kmap["CR1"] == "R30004"
    assert kmap["LT1"] == "DM10100"
    assert kmap["CV1"] == "DM10101"
    assert keyence_address_collisions([row[1] for row in krows[1:]]) == []
    assert keyence_address_collisions(["R30000", "R30000"])
    assert keyence_address_collisions(["R1", "DM1"]) == []

    sysmac = render_tags(rows, target="nj")
    assert "Not a Sysmac Studio import file" in sysmac
    srows = _data_rows(sysmac)
    assert srows[0] == ["name", "data_type", "cio_address", "comment"]
    smap = {row[0]: row for row in srows[1:]}
    assert smap["M1_RUN_FB"][1] == "BOOL"
    assert smap["M1_RUN_FB"][2] == "0.00"
    assert smap["Spare_I0_3"][2] == "0.03"
    assert smap["OL1_OK"][2] == "0.01"
    assert smap["LT1"][1] == "UINT"
    assert smap["LT1"][2] == "2005"
    assert smap["CV1"][2] == "2001"
    assert "Characters that cannot be used together" in sysmac
    assert "Separate underscores stay" in sysmac
    assert omron_address_collisions([row[2] for row in srows[1:]]) == []
    folded = render_tags(
        [
            {"module": "A", "address": "I0_0", "type": "DI", "tag": "M1_RUN_FB", "description": "a"},
            {"module": "A", "address": "I0_1", "type": "DI", "tag": "m1_run_fb", "description": "b"},
            {"module": "A", "address": "I0_2", "type": "DI", "tag": "START", "description": "c"},
            {"module": "A", "address": "I0_3", "type": "DI", "tag": "P_START", "description": "d"},
            {"module": "A", "address": "I0_4", "type": "DI", "tag": "M1__RUN", "description": "e"},
            {"module": "A", "address": "I0_5", "type": "DI", "tag": "_HOLD_", "description": "f"},
        ],
        target="sysmac",
    )
    folded_names = [row[0] for row in _data_rows(folded)[1:]]
    assert folded_names[0] == "M1_RUN_FB"
    assert folded_names[1] == "m1_run_fb_2"
    assert folded_names[2] == "START"
    assert folded_names[3] == "START2"
    assert folded_names[4] == "M1_RUN"
    assert folded_names[5] == "HOLD"
    assert len({name.casefold() for name in folded_names}) == len(folded_names)
    combo = render_tags(
        [
            {"module": "AIO", "address": "AI0", "type": "AI", "tag": "LT1", "description": "level"},
            {"module": "AIO", "address": "AI4", "type": "AI", "tag": "LT5", "description": "extra"},
            {"module": "AIO", "address": "AO0", "type": "AO", "tag": "CV1", "description": "speed"},
            {"module": "AD", "address": "AI0", "type": "AI", "tag": "PT1", "description": "pressure"},
        ],
        target="cxprogrammer",
    )
    combo_rows = {row[0]: row for row in (line.split("\t") for line in combo.splitlines() if line)}
    assert combo_rows["LT1"][2] == "2005"
    assert combo_rows["CV1"][2] == "2001"
    assert combo_rows["LT5"][2] == "unverified"
    assert combo_rows["LT5"][3].startswith("unverified")
    assert combo_rows["PT1"][2] == "2011"

    session = build_plc_motor_io()
    cx_file = session.export_plc_tags(tmp_path / "symbols.csv", target="cxp")
    assert cx_file["verified"] is True
    assert cx_file["format"] == "cxprogrammer-symbol-txt"
    assert cx_file["doc"] == CXPROGRAMMER_DOC
    assert cx_file["path"].endswith(".txt")
    assert "\t" in Path(cx_file["path"]).read_text(encoding="utf-8")
    generic = session.export_plc_tags(tmp_path / "kv.csv", target="keyence")
    assert generic["verified"] is False
    assert generic["format"] == "keyence-generic-csv"
    assert generic["doc"] == KEYENCE_SPECS
    sysmac_file = session.export_plc_tags(tmp_path / "nj.csv", target="sysmac")
    assert sysmac_file["verified"] is False
    assert sysmac_file["doc"] == SYSMAC_DOC
