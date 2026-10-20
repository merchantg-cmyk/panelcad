"""Verified PLC I/O modules: terminal maps, ladder sheets, and I/O-list sync."""

from __future__ import annotations

from math import hypot
from pathlib import Path

import pytest

from panelcad.api.session import Session
from panelcad.cli.agent import Agent
from panelcad.core.electrical import rename_component
from panelcad.core.ink import schematic_ink_rect
from panelcad.core.route import path_hits_obstacles
from panelcad.library.io_modules import find_module, format_point_address, list_modules
from panelcad.render.sheet import parse_pin_ref, pin_world
from panelcad.render.titleblock import working_rect
from panelcad.render.wire_labels import boxes_overlap, mark_box, place_wire_marks
from panelcad.render.wiring import component_obstacle, route_sheet_wires


def _io_sheet(session: Session):
    return next(sheet for sheet in session.project.sheets if sheet.kind == "io" and sheet.drawing_style == "ladder")


def _assert_sheet_geometry(session: Session, sheet_id: str) -> None:
    sheet = session.project.sheet(sheet_id)
    wx, wy, ww, wh = working_rect(sheet)
    x1, y1 = wx + ww, wy + wh
    for comp in sheet.components:
        device = session.catalog.get(comp.library_id)
        left, top, right, bottom = schematic_ink_rect(device, comp.x, comp.y, comp.id, component=comp)
        assert left >= wx - 1.0, (comp.id, left, wx)
        assert top >= wy - 1.0, (comp.id, top, wy)
        assert right <= x1 + 1.0, (comp.id, right, x1)
        assert bottom <= y1 + 1.0, (comp.id, bottom, y1)
    routes = route_sheet_wires(sheet, session.catalog)
    marks = place_wire_marks(sheet, session.catalog, routes, ladder=True)
    for wire in sheet.wires:
        pts = routes[wire.id]
        for point in pts:
            assert wx - 1.0 <= point.x <= x1 + 1.0
            assert wy - 1.0 <= point.y <= y1 + 1.0
        ends: set[str] = set()
        for ref in (wire.frm, wire.to):
            try:
                ends.add(parse_pin_ref(ref)[0])
            except ValueError:
                continue
        obstacles = []
        for comp in sheet.components:
            if comp.id in ends:
                continue
            device = session.catalog.get(comp.library_id)
            if (device.category or "").startswith("connection.notes"):
                continue
            obstacles.append(component_obstacle(sheet, session.catalog, comp, pad=0.2))
        assert not path_hits_obstacles(pts, obstacles), (wire.id, wire.frm, wire.to)
        at = marks.get(wire.id)
        if not at:
            continue
        x, y, align = at
        box = mark_box(x, y, wire.drawn_mark(), align=align, spec=wire.spec_text())
        for start, end in zip(pts, pts[1:]):
            seg = (
                min(start.x, end.x) - 0.45,
                min(start.y, end.y) - 0.45,
                max(start.x, end.x) + 0.45,
                max(start.y, end.y) + 0.45,
            )
            assert not boxes_overlap(box, seg), (wire.id, wire.number)


def _reaches(sheet, start: str, rail: str) -> bool:
    """True when a pin connects to a rail name through junctions."""

    def _match(end: str, token: str) -> bool:
        return end == token or end.split(".", 1)[0] == token

    seen = {start}
    frontier = [start]
    while frontier:
        token = frontier.pop()
        if token == rail:
            return True
        for wire in sheet.wires:
            ends = (wire.frm, wire.to)
            if _match(ends[0], token):
                other = ends[1]
            elif _match(ends[1], token):
                other = ends[0]
            else:
                continue
            head = other.split(".", 1)[0]
            for nxt in (other, head):
                if nxt not in seen:
                    seen.add(nxt)
                    frontier.append(nxt)
    return False


def test_verified_modules_and_refusal():
    session = Session.create("Modules")
    verified = [module.id for module in list_modules(session.catalog) if module.verified]
    assert "brx.bx_08nd3" in verified
    assert "click.c0_16nd3" in verified
    assert "ab.5069_ib16" in verified
    assert "siemens.sm1221_di16" in verified
    with pytest.raises(ValueError, match="will not auto-draw"):
        session.io_add_module("platform.s7_1500")
    with pytest.raises(ValueError, match="will not auto-draw"):
        session.io_add_module("iec.plc.compact_16dio")


def test_terminal_maps_match_the_cited_modules():
    session = Session.create("Maps")
    brx = find_module(session.catalog, "brx.bx_08nd3")
    rows = {row["terminal"]: row for row in brx.terminal_map()}
    assert rows["0"]["common"] == "1C"
    assert rows["3"]["common"] == "1C"
    assert rows["4"]["common"] == "2C"
    assert rows["7"]["common"] == "2C"
    click = find_module(session.catalog, "click.c0_16nd3")
    commons = {row["terminal"]: row["common"] for row in click.terminal_map() if row["role"] == "point"}
    assert commons["1"] == "C1"
    assert commons["4"] == "C1"
    assert commons["13"] == "C4"
    ib = find_module(session.catalog, "ab.5069_ib16")
    assert all(row["role"] != "common" for row in ib.terminal_map())
    assert {row["terminal"] for row in ib.terminal_map() if row["role"] == "nc"} == {"16", "17"}
    ob = find_module(session.catalog, "ab.5069_ob16")
    power = {row["terminal"]: row["role"] for row in ob.terminal_map()}
    assert power["LA+"] == "supply_plus"
    assert power["LA-"] == "supply_minus"
    sm = find_module(session.catalog, "siemens.sm1221_di16")
    sm_rows = {row["terminal"]: row for row in sm.terminal_map()}
    assert sm_rows["X10:4"]["common"] == "1M"
    assert sm_rows["X11:3"]["role"] == "common"


@pytest.mark.confirm_generic_layout
def test_address_syntax_follows_the_platform():
    session = Session.create("Addr")
    cases = [
        ("brx.bx_08nd3", "0", {}, "X0"),
        ("click.c0_16nd3", "1", {"position": 1}, "X101"),
        ("ab.5069_ib16", "I00", {"slot": 1}, "Local:1:I.Pt00.Data"),
        ("ab.5069_ob16", "Q0", {"slot": 2}, "Local:2:O.Pt00.Data"),
        ("ab.5069_if8", "AI0", {"slot": 3}, "Local:3:I.Ch00.Data"),
        ("siemens.sm1221_di16", "Ia0", {"address_base": 0}, "%I0.0"),
        ("ab.2085_iq16", "PT00", {"slot": 1}, "_IO_X1_DI_00"),
    ]
    for module_id, pin, kwargs, expect in cases:
        module = find_module(session.catalog, module_id)
        point = module.point(pin)
        assert format_point_address(module, point, **kwargs) == expect
    added = session.io_add_module("click.c0_16nd3", ref="C1", position=1)
    assert added["points"][0]["address"] == "X101"
    labels = session.project.sheets[-1].component("C1").properties["pin_labels"]
    assert labels["1"] == "X101"
    session.io_generate_sheets()
    rows = {row["address"]: row for row in session.io_rows() if row["module"] == "C1"}
    assert rows["1"]["point"] == "X101"


@pytest.mark.confirm_generic_layout
def test_sink_and_source_commons_and_output_supply():
    sink = Session.create("Sink")
    sink.io_add_module("brx.bx_08nd3", ref="A1", polarity="sink")
    sink.io_generate_sheets()
    sheet = _io_sheet(sink)
    assert (sheet.rail_hot, sheet.rail_return) == ("24VDC", "0V")
    assert _reaches(sheet, "A1.1C", "0V")
    source = Session.create("Source")
    source.io_add_module("brx.bx_08nd3", ref="A1", polarity="source")
    source.io_generate_sheets()
    sheet = _io_sheet(source)
    assert (sheet.rail_hot, sheet.rail_return) == ("0V", "24VDC")
    assert _reaches(sheet, "A1.1C", "24VDC")
    output = Session.create("Out")
    output.io_add_module("ab.5069_ob16", ref="Q1")
    output.io_assign_point("Q1", "Q0", tag="PL1", description="Run", field="pilot_light")
    output.io_generate_sheets(fuses=True)
    sheet = _io_sheet(output)
    assert _reaches(sheet, "Q1.LAplus", "24VDC")
    assert _reaches(sheet, "Q1.LAminus", "0V")
    fuses = [comp for comp in sheet.components if "fuse" in comp.library_id]
    assert fuses and all(comp.rotation == 0 for comp in fuses)
    ib = Session.create("IB")
    ib.io_add_module("ab.5069_ib16", ref="IB1")
    ib.io_generate_sheets()
    sheet = _io_sheet(ib)
    pins = {pin.id for pin in ib.catalog.get(sheet.component("IB1").library_id).pins}
    assert "NC16" not in pins
    note = " ".join(
        str(comp.properties.get("text") or "")
        for comp in sheet.components
        if "notes" in comp.library_id and float(comp.properties.get("width") or 0) > 40
    )
    assert "SA" in note


def _body_rect(session: Session, comp) -> tuple[float, float, float, float]:
    """Symbol or notes rectangle, without the tag band above the symbol."""

    device = session.catalog.get(comp.library_id)
    if (device.category or "").startswith("connection.notes"):
        from panelcad.render.painter import notes_box

        width, height = notes_box(comp, comp.visual_scale())
        return (comp.x + 0.4, comp.y + 0.4, comp.x + width - 0.4, comp.y + height - 0.4)
    left, top, right, bottom = schematic_ink_rect(device, comp.x, comp.y, comp.id, component=comp)
    return (left + 0.3, max(top, comp.y) + 0.3, right - 0.3, bottom - 0.3)


def _assert_wires_miss_bodies(session: Session, sheet_id: str) -> None:
    sheet = session.project.sheet(sheet_id)
    routes = route_sheet_wires(sheet, session.catalog)
    for wire in sheet.wires:
        pts = routes[wire.id]
        ends = set()
        for ref in (wire.frm, wire.to):
            try:
                ends.add(parse_pin_ref(ref)[0])
            except ValueError:
                continue
        for comp in sheet.components:
            if comp.id in ends or "junction" in comp.library_id:
                continue
            rect = _body_rect(session, comp)
            if rect[2] <= rect[0] or rect[3] <= rect[1]:
                continue
            assert not path_hits_obstacles(pts, [rect]), (wire.frm, wire.to, comp.id)


@pytest.mark.confirm_generic_layout
def test_return_wires_miss_the_next_device_and_spare_labels():
    session = Session.create("Clear")
    session.io_add_module("brx.bx_08nd3", ref="A1")
    session.io_assign_point("A1", "0", tag="START", field="pushbutton")
    session.io_assign_point("A1", "1", tag="PX1", field="prox_3wire")
    session.io_assign_point("A1", "2", tag="LS1", field="limit_switch")
    session.io_generate_sheets(fuses=True)
    _assert_wires_miss_bodies(session, _io_sheet(session).id)
    analog = Session.create("Analog")
    analog.io_add_module("ab.5069_if8", ref="AI1", slot=3)
    analog.io_assign_point("AI1", "AI0", tag="TT1", field="transmitter_2wire")
    analog.io_assign_point("AI1", "AI1", tag="PT1", field="transmitter_4wire")
    analog.io_generate_sheets(fuses=True)
    _assert_wires_miss_bodies(analog, _io_sheet(analog).id)
    _assert_sheet_geometry(analog, _io_sheet(analog).id)


@pytest.mark.confirm_generic_layout
def test_sheet_geometry_wire_numbers_and_spare_label():
    session = Session.create("Draw")
    session.io_add_module("brx.bx_08nd3", ref="A1")
    session.io_assign_point("A1", "0", tag="START", description="Start button", field="pushbutton")
    session.io_assign_point("A1", "1", tag="PX1", description="Gate prox", field="prox_3wire")
    session.io_add_module("click.c0_08td1", ref="Y1")
    session.io_assign_point("Y1", "1", tag="SOL1", description="Valve", field="solenoid")
    result = session.io_generate_sheets(fuses=True)
    assert result["sheets"] == ["io-bx-08nd3"]
    sheet = session.project.sheet("io-bx-08nd3")
    assert {comp.id for comp in sheet.components if comp.properties.get("io_module")} == {"A1", "Y1"}
    assert all(wire.number for wire in sheet.wires)
    _assert_sheet_geometry(session, sheet.id)
    rows = {row["address"]: row for row in session.io_rows() if row["module"] == "A1"}
    assert rows["0"]["tag"] == "START"
    assert rows["0"]["point"] == "X0"
    assert rows["0"]["spare"] == "no"
    assert rows["2"]["spare"] == "yes"
    assert rows["2"]["tag"].startswith("SPARE_")
    assert any(comp.properties.get("text") == "SPARE" for comp in sheet.components)
    # A tee (three wire ends) sits on a junction dot.
    routes = route_sheet_wires(sheet, session.catalog)
    ends: dict[tuple[float, float], int] = {}
    for wire in sheet.wires:
        pts = routes[wire.id]
        for point in (pts[0], pts[-1]):
            key = (round(point.x, 1), round(point.y, 1))
            ends[key] = ends.get(key, 0) + 1
    junctions = []
    for comp in sheet.components:
        if "junction" not in comp.library_id:
            continue
        device = session.catalog.get(comp.library_id)
        cx, cy = pin_world(comp, device.pin("E"))
        junctions.append((cx, cy))
    tees = [key for key, count in ends.items() if count >= 3]
    assert tees
    for key in tees:
        assert any(hypot(key[0] - cx, key[1] - cy) <= 1.5 for cx, cy in junctions)
    panel = next(page for page in session.project.sheets if page.kind == "panel_layout")
    assert "xref_io" in panel.component("A1").properties
    assert "io_module" not in panel.component("A1").properties
    xref = sheet.component("A1").properties["xref_rack"]
    assert xref[0].isdigit() and "-" in xref


def test_bom_skips_junctions(tmp_path: Path):
    session = Session.create("BOM")
    session.io_add_module("brx.bx_08nd3", ref="A1")
    session.io_generate_sheets()
    text = session.export_bom(tmp_path / "bom.csv").read_text(encoding="utf-8")
    assert "BX-08ND3" in text
    assert "junction" not in text.lower()


def test_sync_round_trip(tmp_path: Path):
    session = Session.create("Sync")
    session.io_add_module("brx.bx_08nd3", ref="A1")
    session.io_assign_point("A1", "0", tag="START", description="Start", field="pushbutton")
    session.io_generate_sheets()
    sheet = session.project.sheet("io-bx-08nd3")
    rename_component(session.project, sheet, "START", "RUN")
    sheet.component("RUN").properties["io_description"] = "Running"
    rows = {row["address"]: row for row in session.io_rows() if row["module"] == "A1"}
    assert rows["0"]["tag"] == "RUN"
    assert rows["0"]["description"] == "Running"
    assert rows["0"]["device"] == "RUN"
    src = tmp_path / "points.csv"
    src.write_text(
        "module,address,description,tag,field\nA1,1,Gate,LS1,limit_switch\n",
        encoding="utf-8",
        newline="",
    )
    imported = session.import_io_list(src)
    assert imported["errors"] == []
    assert "A1.1" in imported["updated"]
    assert sheet.component("A1").properties["pin_labels"]["1"] == "X1"
    result = session.io_generate_sheets()
    assert result["warnings"] == [] or all("SPARE" not in warning for warning in result["warnings"])
    sheet = session.project.sheet("io-bx-08nd3")
    assert sheet.component("RUN").library_id.endswith("pb_no")
    assert "limit" in sheet.component("LS1").library_id
    session.io_assign_point("A1", "2", tag="BAD", description="coil", field="relay_coil")
    result = session.io_generate_sheets()
    assert any("SPARE" in warning for warning in result["warnings"])
    rows = {row["address"]: row for row in session.io_rows() if row["module"] == "A1"}
    assert rows["2"]["spare"] == "yes"
    assert sheet.component("RUN").library_id.endswith("pb_no")


def test_agent_and_cli(tmp_path: Path):
    agent = Agent()
    assert agent.handle({"id": 1, "op": "create_project", "name": "Agent IO"})["ok"]
    added = agent.handle({"id": 2, "op": "io_add_module", "module": "brx.bx_08nd3", "ref": "A1"})
    assert added["ok"], added
    assigned = agent.handle(
        {"id": 3, "op": "io_assign_point", "module": "A1", "pin": "0", "tag": "START", "field": "pushbutton"}
    )
    assert assigned["ok"], assigned
    drawn = agent.handle({"id": 4, "op": "io_generate_sheets"})
    assert drawn["ok"], drawn
    assert drawn["result"]["sheets"]
    refused = agent.handle({"id": 5, "op": "io_add_module", "module": "platform.productivity"})
    assert not refused["ok"]
    assert "will not auto-draw" in refused["error"]
    help_doc = agent.handle({"id": 6, "op": "help"})
    assert "io_generate_sheets" in help_doc["result"]["schemas"]

    from typer.testing import CliRunner

    from panelcad.cli.main import app

    project = tmp_path / "io.pcad.yaml"
    Session.create("CLI").save(project)
    runner = CliRunner()
    listed = runner.invoke(app, ["io", "modules", "--platform", "click"])
    assert listed.exit_code == 0, listed.output
    assert "C0-16ND3" in listed.output
    assert "unverified" in listed.output
    generated = runner.invoke(
        app,
        ["io", "generate", str(project), "--module", "brx.bx_12nd3", "--ref", "A1", "--fuses"],
    )
    assert generated.exit_code == 0, generated.output
    opened = Session.open(project)
    assert any(comp.id == "A1" for sheet in opened.project.sheets for comp in sheet.components)


def test_wizard_generates_a_preview():
    from PySide6.QtCore import Qt
    from PySide6.QtWidgets import QApplication

    from panelcad.gui.io_wizard import IoWizard

    if QApplication.instance() is None:
        QApplication([])
    session = Session.create("Wizard")
    dialog = IoWizard(session)
    checked = False
    for row in range(dialog.modules.count()):
        item = dialog.modules.item(row)
        if item.text().startswith("BX-08ND3"):
            item.setCheckState(Qt.CheckState.Checked)
            checked = True
    assert checked
    dialog._fill_points()
    assert dialog.points.rowCount() == 8
    result = dialog.generate()
    assert result["sheets"]
    dialog._paint_preview()
    assert dialog.preview.pixmap() is not None and not dialog.preview.pixmap().isNull()


def test_wizard_csv_keeps_field_codes(tmp_path):
    from PySide6.QtCore import Qt
    from PySide6.QtWidgets import QApplication

    from panelcad.gui.io_wizard import IoWizard

    if QApplication.instance() is None:
        QApplication([])
    dialog = IoWizard(Session.create("Field codes"))
    checked = False
    for row in range(dialog.modules.count()):
        item = dialog.modules.item(row)
        if item.text().startswith("BX-08ND3"):
            item.setCheckState(Qt.CheckState.Checked)
            checked = True
            assert "di_sink" not in item.text()
            assert "8 inputs, sink/source" in item.text()
    assert checked
    dialog._fill_points()
    module = dialog.points.item(0, 0).text()
    pin_limit = dialog.points.item(0, 1).text()
    pin_prox = dialog.points.item(1, 1).text()
    path = tmp_path / "points.csv"
    path.write_text(
        "Module,Pin,Tag,Description,Field\n"
        f"{module},{pin_limit},LS1,Door limit,limit_switch\n"
        f"{module},{pin_prox},PX1,Prox,prox_3wire\n",
        encoding="utf-8",
    )
    assert dialog.load_points_csv(path) == 2
    fields = {row["tag"]: row["field"] for row in dialog.assignments()}
    assert fields["LS1"] == "limit_switch"
    assert fields["PX1"] == "prox_3wire"
    label = tmp_path / "label.csv"
    label.write_text(
        "Module,Pin,Tag,Description,Field\n"
        f"{module},{pin_limit},LS2,Door,Limit switch\n",
        encoding="utf-8",
    )
    dialog.load_points_csv(label)
    labeled = {row["tag"]: row["field"] for row in dialog.assignments()}
    assert labeled["LS2"] == "limit_switch"


@pytest.mark.confirm_generic_layout
def test_separate_rails_do_not_share_a_sheet():
    session = Session.create("Rails")
    session.io_add_module("brx.bx_08nd3", ref="A1", polarity="source")
    session.io_add_module("click.c0_08td1", ref="Y1")
    result = session.io_generate_sheets()
    assert len(result["sheets"]) == 2


@pytest.mark.confirm_generic_layout
def test_second_generate_leaves_the_rack_unchanged():
    import json

    session = Session.create("Twice")
    session.io_add_module("brx.bx_08nd3", ref="A1")
    session.io_generate_sheets()
    panel = next(sheet for sheet in session.project.sheets if sheet.kind == "panel_layout")
    before = json.dumps(panel.model_dump(mode="json"), sort_keys=True)
    again = session.io_generate_sheets()
    after = json.dumps(panel.model_dump(mode="json"), sort_keys=True)
    assert before == after
    assert panel.id not in again["sheets"]
    assert panel.wires == []
    foot = panel.component("A1")
    assert "io_module" not in foot.properties
    assert foot.properties["io_module_ref"] == "brx.bx_08nd3"
    session.io_add_module("ab.5069_ib16", ref="IB1")
    mixed = session.io_generate_sheets()
    assert panel.id not in mixed["sheets"]
    assert panel.wires == []
    assert all("io_module" not in comp.properties for comp in panel.components)
    # 5069-IB16 has no published body — stay off the rack until a catalog
    # footprint or a confirmed envelope exists.
    assert all(comp.id != "IB1" for comp in panel.components)


def test_module_refs_are_unique_across_sheets():
    session = Session.create("Refs")
    first = session.io_add_module("brx.bx_08nd3", polarity="source")
    second = session.io_add_module("click.c0_08td1")
    assert first["sheet"] != second["sheet"]
    assert first["ref"] == "IO1"
    assert second["ref"] == "IO2"
    named = session.io_add_module("ab.5069_ob16", ref="Q1")
    assert named["ref"] == "Q1"

    agent = Agent()
    assert agent.handle({"id": 1, "op": "create_project", "name": "Refs"})["ok"]
    one = agent.handle({"id": 2, "op": "io_add_module", "module": "brx.bx_08nd3", "polarity": "source"})
    two = agent.handle({"id": 3, "op": "io_add_module", "module": "click.c0_08td1"})
    assert one["ok"] and two["ok"]
    assert one["result"]["ref"] == "IO1"
    assert two["result"]["ref"] == "IO2"


@pytest.mark.confirm_generic_layout
def test_sheet_address_is_what_the_exporters_write(tmp_path: Path):
    from panelcad.core.checks import run_checks

    click = Session.create("Click")
    click.io_add_module("click.c0_08td1", ref="Y1", position=2)
    click.io_generate_sheets()
    points = {row["point"] for row in click.io_rows() if row["module"] == "Y1"}
    assert "Y201" in points
    with pytest.raises(ValueError, match="does not match"):
        click.export_plc_tags(tmp_path / "click-5000.csv", target="studio5000")
    with pytest.raises(ValueError, match="does not match"):
        click.export_plc_tags(tmp_path / "click-domore.csv", target="domore")

    cards = Session.create("Cards")
    cards.io_add_module("brx.bx_08nd3", ref="A1")
    cards.io_add_module("brx.bx_08nd3", ref="A2")
    cards.io_generate_sheets()
    by_module: dict[str, set[str]] = {}
    for row in cards.io_rows():
        if row["type"] == "DI":
            by_module.setdefault(row["module"], set()).add(row["point"])
    assert "X0" in by_module["A1"]
    assert "X0" not in by_module["A2"]
    assert "X8" in by_module["A2"]
    cards.export_plc_tags(tmp_path / "brx.csv", target="domore")
    exported = (tmp_path / "brx.csv").read_text(encoding="utf-8")
    assert "X0" in exported and "X8" in exported

    forced = Session.create("Dup")
    forced.io_add_module("brx.bx_08nd3", ref="A1", address_base=0)
    forced.io_add_module("brx.bx_08nd3", ref="A2", address_base=0)
    forced.io_generate_sheets()
    assert any(item.check == "duplicate_address" for item in run_checks(forced.project, forced.catalog))

    sm = Session.create("SM")
    sm.io_add_module("siemens.sm1221_di16", ref="A1", address_base=8)
    sm.io_generate_sheets()
    assert any(row["point"] == "%I8.0" for row in sm.io_rows() if row["module"] == "A1")
    sm.export_plc_tags(tmp_path / "sm.csv", target="tia_csv")
    assert "%I8.0" in (tmp_path / "sm.csv").read_text(encoding="utf-8")
    label = next(
        comp.pin_overrides["1M"].label
        for sheet in sm.project.sheets
        for comp in sheet.components
        if comp.id == "A1" and comp.properties.get("io_module")
    )
    assert "1M" in label and "X10:3" in label

    micro = Session.create("M8")
    micro.io_add_module("ab.2085_iq16", ref="A1", slot=2)
    micro.io_generate_sheets()
    assert any(row["point"] == "_IO_X2_DI_00" for row in micro.io_rows())
    with pytest.raises(ValueError, match="does not match"):
        micro.export_plc_tags(tmp_path / "m8.csv", target="studio5000")

    ob = Session.create("OB")
    ob.io_add_module("ab.5069_ob16", ref="Q1", slot=2)
    ob.io_generate_sheets()
    assert any(row["point"] == "Local:2:O.Pt00.Data" for row in ob.io_rows())
    ob.export_plc_tags(tmp_path / "ob.csv", target="studio5000")
    assert "Local:2:O.Pt00.Data" in (tmp_path / "ob.csv").read_text(encoding="utf-8")


@pytest.mark.confirm_generic_layout
def test_if8_shield_and_four_wire_do_not_bond_to_0v():
    session = Session.create("IF8")
    session.io_add_module("ab.5069_if8", ref="AI1", slot=1)
    session.io_assign_point("AI1", "AI0", tag="TT1", field="transmitter_4wire")
    session.io_generate_sheets()
    sheet = _io_sheet(session)
    card = sheet.component("AI1")
    assert "SH16" in card.pin_overrides
    assert "SH17" in card.pin_overrides
    assert any(comp.library_id.endswith("connection.ground") for comp in sheet.components)
    assert _reaches(sheet, "AI1.M0", "TT1")
    assert not _reaches(sheet, "AI1.M0", "0V")
    assert not _reaches(sheet, "AI1.SH16", "0V")


@pytest.mark.confirm_generic_layout
def test_click_output_plus_v_reaches_the_hot_rail():
    session = Session.create("Plus")
    session.io_add_module("click.c0_08td1", ref="Y1")
    session.io_generate_sheets(fuses=True)
    sheet = _io_sheet(session)
    assert _reaches(sheet, "Y1.VP", "24VDC")


@pytest.mark.confirm_generic_layout
def test_placeholder_footprint_and_rtb_are_on_the_bom(tmp_path: Path):
    from panelcad.core.checks import run_checks

    session = Session.create("PH")
    session.io_add_module("ab.5069_ib16", ref="IB1")
    session.io_generate_sheets()
    text = session.export_bom(tmp_path / "bom.csv").read_text(encoding="utf-8")
    assert "5069-RTB18-SCREW" in text
    assert "PLACEHOLDER" in text
    assert any(item.check == "placeholder_footprint" for item in run_checks(session.project, session.catalog))
    card = next(
        comp
        for sheet in session.project.sheets
        for comp in sheet.components
        if comp.id == "IB1" and comp.properties.get("io_module")
    )
    assert card.properties.get("placeholder_footprint") is True
    assert all(
        comp.id != "IB1"
        for sheet in session.project.sheets
        if sheet.kind == "panel_layout"
        for comp in sheet.components
    )


def test_source_input_fuses_the_positive_rail():
    session = Session.create("SrcFuse")
    session.io_add_module("brx.bx_08nd3", ref="A1", polarity="source")
    session.io_generate_sheets(fuses=True)
    sheet = _io_sheet(session)
    fuses = [comp for comp in sheet.components if "fuse" in comp.library_id]
    assert fuses
    for fuse in fuses:
        assert _reaches(sheet, fuse.id, "24VDC")
        assert not _reaches(sheet, fuse.id, "0V")


@pytest.mark.confirm_generic_layout
def test_sixteen_points_continue_on_a2():
    session = Session.create("Cont")
    session.io_add_module("ab.5069_ib16", ref="IB1")
    result = session.io_generate_sheets()
    assert len(result["sheets"]) == 2
    notes = []
    for sheet_id in result["sheets"]:
        sheet = session.project.sheet(sheet_id)
        assert sheet.page.size == "A2"
        assert sheet.page.orientation == "landscape"
        assert sheet.page.width_mm is None
        for comp in sheet.components:
            notes.append(str(comp.properties.get("text") or ""))
    blob = "\n".join(notes)
    assert "CONT. ON SHEET" in blob
    assert "CONT. FROM SHEET" in blob
    assert any(comp.properties.get("io_continuation") == "IB1" for sheet in session.project.sheets for comp in sheet.components)
    assert any(comp.properties.get("shown_tag") == "IB1" for sheet in session.project.sheets for comp in sheet.components)


@pytest.mark.confirm_generic_layout
def test_csv_field_round_trip_and_terminal_token(tmp_path: Path):
    session = Session.create("CSV")
    session.io_add_module("brx.bx_08nd3", ref="A1")
    session.io_assign_point("A1", "0", tag="LS1", field="limit_switch")
    session.io_generate_sheets()
    path = session.export_io_list(tmp_path / "io.csv")
    text = path.read_text(encoding="utf-8")
    assert "limit_switch" in text
    edited = text.replace("limit_switch", "prox_3wire")
    path.write_text(edited, encoding="utf-8")
    imported = session.import_io_list(path)
    assert imported["errors"] == []
    sheet = session.project.sheet("io-bx-08nd3")
    assert sheet.component("A1").properties["io"]["0"]["field"] == "prox_3wire"
    # Import redraws. The list and the symbol stay the same device.
    assert "prox" in sheet.component("LS1").library_id

    from typer.testing import CliRunner

    from panelcad.cli.main import app

    project = tmp_path / "ib.pcad.yaml"
    points = tmp_path / "points.csv"
    points.write_text("pin,tag,description,field\n0,START,Run,pushbutton\n", encoding="utf-8")
    Session.create("IB").save(project)
    runner = CliRunner()
    ok = runner.invoke(app, ["io", "generate", str(project), "--module", "ab.5069_ib16", "--ref", "IB1", "--points", str(points)])
    assert ok.exit_code == 0, ok.output
    assert "Traceback" not in ok.output
    bad = tmp_path / "bad.csv"
    bad.write_text("pin,tag\nNO-SUCH,X\n", encoding="utf-8")
    failed = runner.invoke(
        app, ["io", "generate", str(tmp_path / "bad.pcad.yaml"), "--module", "ab.5069_ib16", "--points", str(bad)]
    )
    assert failed.exit_code != 0
    assert "Traceback" not in (failed.output + failed.stderr)


@pytest.mark.confirm_generic_layout
def test_spare_tags_are_unique_and_a_fresh_sheet_is_quiet():
    from panelcad.core.checks import run_checks

    session = Session.create("Quiet")
    session.io_add_module("brx.bx_08nd3", ref="A1")
    session.io_add_module("brx.bx_08nd3", ref="A2")
    session.io_generate_sheets()
    tags = [row["tag"] for row in session.io_rows() if row["spare"] == "yes"]
    assert tags
    assert len(tags) == len(set(tags))
    noisy = [
        item
        for item in run_checks(session.project, session.catalog)
        if item.severity in {"error", "warning"} and not str(item.check).startswith("nameplate_")
    ]
    assert not noisy, [(item.check, item.message) for item in noisy]


@pytest.mark.confirm_generic_layout
def test_second_module_shares_the_sheet_after_the_first_is_drawn():
    session = Session.create("Share")
    session.io_add_module("brx.bx_08nd3", ref="A1")
    session.io_generate_sheets()
    session.io_add_module("click.c0_08td1", ref="Y1")
    session.io_assign_point("Y1", "1", tag="SOL1", field="solenoid")
    session.io_assign_point("Y1", "2", tag="PL1", field="pilot_light")
    result = session.io_generate_sheets(fuses=True)
    assert result["sheets"] == ["io-bx-08nd3"]
    sheet = session.project.sheet("io-bx-08nd3")
    ids = {comp.id for comp in sheet.components}
    assert "A1" in ids and "Y1" in ids
    assert _reaches(sheet, "SOL1", "24VDC")
    assert _reaches(sheet, "Y1.VP", "24VDC")
    # A dot only where three wires meet.
    counts: dict[str, int] = {}
    for wire in sheet.wires:
        for ref in (wire.frm, wire.to):
            head = ref.split(".", 1)[0]
            counts[head] = counts.get(head, 0) + 1
    for comp in sheet.components:
        if "junction" not in comp.library_id:
            continue
        assert counts.get(comp.id, 0) >= 3, comp.id


def test_delete_field_device_stays_spare_and_rename_reaches_the_list():
    session = Session.create("Edit")
    session.io_add_module("brx.bx_08nd3", ref="A1")
    session.io_assign_point("A1", "0", tag="START", description="Start", field="pushbutton")
    session.io_generate_sheets()
    session.delete_component("START")
    session.io_generate_sheets()
    sheet = session.project.sheet("io-bx-08nd3")
    assert "START" not in {comp.id for comp in sheet.components}
    meta = sheet.component("A1").properties["io"]["0"]
    assert meta["field"] == "spare"
    assert str(meta["tag"]).startswith("SPARE_")
    session.io_assign_point("A1", "1", tag="LS1", description="Gate", field="limit_switch")
    session.io_generate_sheets()
    session.rename_component("LS1", "GATE")
    rows = {row["address"]: row for row in session.io_rows() if row["module"] == "A1"}
    assert rows["1"]["tag"] == "GATE"


def _conducts(comp) -> bool:
    lid = comp.library_id.lower()
    role = str(comp.properties.get("io_role") or "")
    return any(token in lid for token in ("junction", "waypoint", "ground", "fuse", "terminal")) or role.startswith(
        "fuse"
    )


def _net(session: Session, sheet, start: str) -> set[str]:
    """Pins and rails reachable from ``start``. Field devices and cards do not pass the net through."""

    by_id = {comp.id: comp for comp in sheet.components}
    seen = {start}
    frontier = [start]
    while frontier:
        token = frontier.pop()
        for wire in sheet.wires:
            ends = (wire.frm, wire.to)
            if token not in ends:
                continue
            other = ends[1] if ends[0] == token else ends[0]
            if other not in seen:
                seen.add(other)
                frontier.append(other)
            head = other.split(".", 1)[0]
            comp = by_id.get(head)
            if comp is None or not _conducts(comp):
                continue
            device = session.catalog.get(comp.library_id)
            for pin in device.pins:
                ref = f"{head}.{pin.id}"
                if ref not in seen:
                    seen.add(ref)
                    frontier.append(ref)
    return seen


def _populated() -> Session:
    session = Session.create("Review")
    specs = [
        ("brx.bx_08nd3", "IO1", {}),
        ("click.c0_08td1", "IO2", {}),
        ("ab.5069_ib16", "IO3", {"slot": 1}),
        ("ab.5069_ob16", "IO4", {"slot": 2}),
        ("ab.5069_if8", "IO5", {"slot": 3}),
        ("siemens.sm1221_di16", "IO6", {"address_base": 0}),
        ("brx.bx_12nd3", "IO7", {"polarity": "source"}),
        ("ab.2085_iq16", "IO8", {"slot": 1}),
        ("click.c0_16nd3", "IO9", {"position": 2}),
        ("brx.bx_08nd3", "IO10", {}),
    ]
    for module, ref, extra in specs:
        session.io_add_module(module, ref=ref, **extra)
    session.io_assign_point("IO1", "0", tag="START", description="Start pushbutton", field="pushbutton")
    session.io_assign_point("IO1", "1", tag="PX1", description="Gate prox", field="prox_3wire")
    session.io_assign_point("IO2", "1", tag="SOL1", description="Valve", field="solenoid")
    session.io_assign_point("IO3", "I00", tag="IB_00", description="Prox 0", field="prox_3wire")
    session.io_assign_point("IO3", "I01", tag="IB_01", description="Prox 1", field="prox_3wire")
    session.io_assign_point("IO3", "I02", tag="IB_02", description="Prox 2", field="prox_3wire")
    session.io_assign_point("IO4", "Q0", tag="CR_FWD", description="Forward", field="relay_coil")
    session.io_assign_point("IO4", "Q1", tag="CR_RUN", description="Run", field="relay_coil")
    session.io_assign_point("IO5", "AI0", tag="TT1", description="Level", field="transmitter_2wire")
    session.io_assign_point("IO5", "AI1", tag="TT2", description="Pressure", field="transmitter_4wire")
    session.io_assign_point("IO7", "0", tag="PXS", description="Source prox", field="prox_3wire")
    session.io_assign_point("IO7", "1", tag="PXS2", description="Source prox 2", field="prox_3wire")
    session.io_generate_sheets(fuses=True)
    return session


def _assert_io_drawing_clear(session: Session) -> None:
    """No different-net crossings, and no tag, caption, or wire number on a symbol or another net."""

    from panelcad.core.ink import text_box
    from panelcad.core.pins import io_terminal_caption
    from panelcad.render.hops import (
        _proper_cross,
        _same_net,
        conductor_hits,
        label_routes,
        rendered_wire_paths,
    )
    from panelcad.render.painter import notes_box
    from panelcad.render.sheet_labels import caption_layouts, place_device_tags
    from panelcad.render.wire_labels import mark_box, place_wire_marks

    for sheet in session.project.sheets:
        if sheet.kind != "io":
            continue
        assert conductor_hits(sheet, session.catalog) == []
        routes = route_sheet_wires(sheet, session.catalog)
        wires = [wire for wire in sheet.wires if len(routes.get(wire.id) or []) == 2]
        for index, left in enumerate(wires):
            for right in wires[index + 1 :]:
                if not _same_net(left, right):
                    continue
                cross = _proper_cross(
                    routes[left.id][0], routes[left.id][1], routes[right.id][0], routes[right.id][1]
                )
                assert cross is None, (sheet.id, left.number, cross)
        paths = rendered_wire_paths(sheet, session.catalog, routes)
        marks = place_wire_marks(sheet, session.catalog, label_routes(routes, paths), ladder=True)
        tags = place_device_tags(sheet, session.catalog, ladder=True, project=session.project)

        def _segs(wire_id: str, skip: set[str]) -> list[tuple]:
            found = []
            for wire in sheet.wires:
                if wire.id == wire_id or wire.id in skip:
                    continue
                for sub in paths.get(wire.id) or []:
                    found.extend(zip(sub, sub[1:]))
            return found

        for comp in sheet.components:
            device = session.catalog.get(comp.library_id)
            if device is None:
                continue
            boxes = []
            if (device.category or "").lower().startswith("connection.notes"):
                width, height = notes_box(comp, comp.visual_scale())
                boxes.append((comp.x, comp.y, comp.x + width, comp.y + height))
            elif (
                not comp.properties.get("ladder_card")
                and "junction" not in comp.library_id
                and "waypoint" not in comp.library_id
            ):
                if not comp.properties.get("hide_tag"):
                    layout = tags.get(comp.id)
                    if layout is not None:
                        box = text_box(layout[0], layout[1], comp.id, layout[3], layout[2])
                        if box is not None:
                            boxes.append(box)
                    for cx, cy, align, size, line in caption_layouts(comp, tags.get(comp.id)):
                        box = text_box(cx, cy, line, size, align)
                        if box is not None:
                            boxes.append(box)
                caption = io_terminal_caption(
                    comp, device.schematic_width, device.schematic_height, comp.visual_scale()
                )
                if caption is not None:
                    box = text_box(caption[0], caption[1], caption[4], caption[3], caption[2])
                    if box is not None:
                        boxes.append(box)
            if (device.category or "").lower().startswith("connection.notes"):
                width, height = notes_box(comp, comp.visual_scale())
                note_box = (comp.x, comp.y, comp.x + width, comp.y + height)
                for other in sheet.components:
                    if (
                        other is comp
                        or "junction" in other.library_id
                        or "waypoint" in other.library_id
                    ):
                        continue
                    body = component_obstacle(sheet, session.catalog, other, pad=0.0)
                    overlap_w = min(note_box[2], body[2]) - max(note_box[0], body[0])
                    overlap_h = min(note_box[3], body[3]) - max(note_box[1], body[1])
                    assert overlap_w <= 0.6 or overlap_h <= 0.6, (sheet.id, comp.id, other.id)
            own: set[str] = set()
            for wire in sheet.wires:
                for ref in (wire.frm, wire.to):
                    try:
                        if parse_pin_ref(ref)[0] == comp.id:
                            own.add(wire.id)
                    except ValueError:
                        pass
            foreign = _segs("", own)
            for box in boxes:
                for start, end in foreign:
                    seg = (
                        min(start.x, end.x) - 0.3,
                        min(start.y, end.y) - 0.3,
                        max(start.x, end.x) + 0.3,
                        max(start.y, end.y) + 0.3,
                    )
                    assert not boxes_overlap(box, seg), (sheet.id, comp.id)
        for wire in sheet.wires:
            at = marks.get(wire.id)
            if not at:
                continue
            box = mark_box(at[0], at[1], wire.drawn_mark(), align=at[2], spec=wire.spec_text())
            for other in sheet.wires:
                if other.id == wire.id or (other.number and other.number == wire.number):
                    continue
                for sub in paths.get(other.id) or []:
                    for start, end in zip(sub, sub[1:]):
                        seg = (
                            min(start.x, end.x) - 0.15,
                            min(start.y, end.y) - 0.15,
                            max(start.x, end.x) + 0.15,
                            max(start.y, end.y) + 0.15,
                        )
                        assert not boxes_overlap(box, seg), (sheet.id, wire.number, other.number)
            for comp in sheet.components:
                if "junction" in comp.library_id or "waypoint" in comp.library_id:
                    continue
                if comp.properties.get("ladder_card"):
                    continue
                ends = set()
                for ref in (wire.frm, wire.to):
                    try:
                        ends.add(parse_pin_ref(ref)[0])
                    except ValueError:
                        pass
                if comp.id in ends:
                    continue
                body = component_obstacle(sheet, session.catalog, comp, pad=0.05)
                assert not boxes_overlap(box, body), (sheet.id, wire.number, comp.id)


@pytest.mark.confirm_generic_layout
def test_fuse_refs_are_unique_and_a_second_generate_matches(tmp_path: Path):
    session = _populated()
    fuses = [
        comp.id
        for sheet in session.project.sheets
        for comp in sheet.components
        if "fuse" in comp.library_id
    ]
    assert len(fuses) >= 3
    assert len(fuses) == len(set(fuses))
    from panelcad.core.checks import run_checks

    checks = run_checks(session.project, session.catalog)
    assert not any(item.check == "duplicate_device_tag" for item in checks)

    first = {
        sheet.id: session.export_svg(tmp_path / "a" / f"{sheet.id}.svg", sheet_id=sheet.id).read_bytes()
        for sheet in session.project.sheets
    }
    session.io_generate_sheets(fuses=True)
    second = {
        sheet.id: session.export_svg(tmp_path / "b" / f"{sheet.id}.svg", sheet_id=sheet.id).read_bytes()
        for sheet in session.project.sheets
    }
    assert first.keys() == second.keys()
    for sheet_id, blob in first.items():
        assert blob == second[sheet_id], sheet_id


@pytest.mark.confirm_generic_layout
def test_two_wire_transmitter_returns_on_the_minus_terminal():
    session = Session.create("Loop")
    session.io_add_module("ab.5069_if8", ref="AI1", slot=3)
    session.io_assign_point("AI1", "AI0", tag="TT1", field="transmitter_2wire")
    session.io_generate_sheets(fuses=True)
    sheet = _io_sheet(session)
    assert any(str(comp.properties.get("text") or "") == "EXT SUPPLY" for comp in sheet.components)
    plus = _net(session, sheet, "TT1.PLUS")
    assert "24VDC" not in plus
    assert "0V" not in plus
    minus = _net(session, sheet, "AI1.M0")
    assert "AI1.M0" in minus
    assert "24VDC" not in minus
    assert "0V" not in minus
    assert any(
        "waypoint" in sheet.component(token.split(".", 1)[0]).library_id for token in minus if "." in token
    )
    signal = _net(session, sheet, "TT1.SIG")
    assert "AI1.AI0" in signal
    assert "24VDC" not in signal
    assert "0V" not in signal


@pytest.mark.confirm_generic_layout
def test_rack_modules_stay_inside_and_point_at_their_zone():
    from panelcad.core.checks import run_checks
    from panelcad.core.layout import footprint_size
    from panelcad.core.zones import sheet_number, zone_at

    session = _populated()
    checks = run_checks(session.project, session.catalog)
    assert not any(item.check == "outside_enclosure" for item in checks)
    racks = [sheet for sheet in session.project.sheets if sheet.kind == "panel_layout"]
    modules = [
        comp
        for rack in racks
        for comp in rack.components
        if comp.properties.get("io_module_ref") or comp.properties.get("io_module")
    ]
    # Catalog-dimension modules only. Placeholder CompactLogix / Click / Micro800
    # cards stay on the I/O sheets until someone confirms an envelope.
    on_rack = {comp.id for comp in modules}
    assert {"IO1", "IO6", "IO7", "IO10"} <= on_rack
    assert on_rack.isdisjoint({"IO2", "IO3", "IO4", "IO5", "IO8", "IO9"})
    for rack in racks:
        enc = rack.enclosure
        assert enc is not None
        for comp in rack.components:
            if not (comp.properties.get("io_module_ref") or comp.properties.get("io_module")):
                continue
            device = session.catalog.get(comp.library_id)
            width, height = footprint_size(device, comp)
            assert enc.x - 0.1 <= comp.x
            assert comp.x + width <= enc.x + enc.width_mm + 0.1
            assert enc.y - 0.1 <= comp.y
            assert comp.y + height <= enc.y + enc.height_mm + 0.1
            row, col, _index = zone_at(rack, comp.x + width / 2.0, comp.y + min(height, 30.0) / 2.0)
            number = sheet_number(session.project.sheets, rack.id) or 1
            home = next(
                card
                for page in session.project.sheets
                if page.kind == "io"
                for card in page.components
                if card.id == comp.id and card.properties.get("io_module")
            )
            assert home.properties["xref_rack"] == f"{number}-{row}{col}"


@pytest.mark.confirm_generic_layout
def test_drawing_cleanup_addresses_the_review(tmp_path: Path):
    session = _populated()
    _assert_io_drawing_clear(session)
    paired = session.project.sheet("io-bx-08nd3")
    svg = session.export_svg(tmp_path / "paired.svg", sheet_id=paired.id).read_text(encoding="utf-8")
    dots = [line for line in svg.splitlines() if 'r="1.300"' in line]
    assert dots
    assert all('fill="#222222"' in line for line in dots)

    cont = next(sheet for sheet in session.project.sheets if sheet.io_cont)
    assert "io-cont" not in sheet_text(cont)
    assert "points 8–15" in cont.description
    assert "(cont.)" in cont.description

    ob = next(sheet for sheet in session.project.sheets if "ob16" in sheet.id and not sheet.io_cont)
    labels = [
        str(ov.label)
        for comp in ob.components
        if comp.properties.get("io_module")
        for ov in comp.pin_overrides.values()
    ]
    assert "LA+" in labels
    assert "LA-" in labels
    assert not any("LAplus" in label or "LAminus" in label for label in labels)
    coils = [comp for comp in ob.components if comp.library_id.endswith("contactors.coil")]
    assert coils
    assert {comp.properties.get("coil") for comp in coils} == {"24VDC"}
    from panelcad.core.checks import run_checks

    assert not any(item.check == "coil_voltage_mismatch" for item in run_checks(session.project, session.catalog))

    sm = find_module(session.catalog, "siemens.sm1221_di16")
    panel = session.catalog.get(sm.library_id).panel
    assert panel.dim_source == "catalog"
    assert (panel.width_mm, panel.height_mm, panel.depth_mm) == (45.0, 100.0, 75.0)
    footprint = next(
        comp
        for sheet in session.project.sheets
        if sheet.kind == "panel_layout"
        for comp in sheet.components
        if comp.id == "IO6"
    )
    assert footprint.properties.get("placeholder_footprint") is False
    sm_sheet = next(
        sheet for sheet in session.project.sheets if "6es7221" in sheet.id and not sheet.io_cont
    )
    sm_card = next(comp for comp in sm_sheet.components if comp.properties.get("io_module"))
    common = sm_card.pin_overrides["1M"].label
    assert "1M" in common and "X10:3" in common
    if8 = next(sheet for sheet in session.project.sheets if sheet.id == "io-5069-if8")
    assert sum(1 for comp in if8.components if comp.properties.get("io_module")) == 1


def sheet_text(sheet) -> str:
    return f"{sheet.description} {sheet.io_cont}"


def test_cli_redraw_rename_and_delete(tmp_path: Path):
    from typer.testing import CliRunner

    from panelcad.cli.main import app

    session = Session.create("CLI2")
    session.io_add_module("brx.bx_08nd3", ref="A1")
    session.io_assign_point("A1", "0", tag="START", description="Start", field="pushbutton")
    session.io_generate_sheets()
    project = tmp_path / "io.pcad.yaml"
    session.save(project)
    runner = CliRunner()
    missing = runner.invoke(app, ["io", "generate", str(project)])
    assert missing.exit_code == 1
    assert "redraw" in (missing.stderr or missing.output)
    redrawn = runner.invoke(app, ["io", "regenerate", str(project)])
    assert redrawn.exit_code == 0, redrawn.output
    again = runner.invoke(app, ["io", "generate", str(project), "--all"])
    assert again.exit_code == 0, again.output
    opened = Session.open(project)
    assert opened.project.sheet("io-bx-08nd3").component("START").library_id.endswith("pb_no")
    renamed = runner.invoke(app, ["rename", str(project), "START", "RUN"])
    assert renamed.exit_code == 0, renamed.output
    assert Session.open(project).project.sheet("io-bx-08nd3").component("RUN")
    deleted = runner.invoke(app, ["delete", str(project), "RUN"])
    assert deleted.exit_code == 0, deleted.output
    assert "RUN" not in {
        comp.id for comp in Session.open(project).project.sheet("io-bx-08nd3").components
    }


def test_gui_regenerate_and_refused_tag_export(tmp_path: Path, monkeypatch: pytest.MonkeyPatch):
    from PySide6.QtWidgets import QApplication, QMessageBox

    from panelcad.gui.main_window import MainWindow
    from panelcad.gui.plc_tag_dialog import PlcTagExportDialog

    if QApplication.instance() is None:
        QApplication([])
    session = Session.create("Gui")
    session.io_add_module("brx.bx_08nd3", ref="A1")
    session.io_assign_point("A1", "0", tag="START", field="pushbutton")
    session.io_generate_sheets()
    window = MainWindow(session)
    reports = next(
        action.menu()
        for action in window.menuBar().actions()
        if action.text().replace("&", "") == "Reports"
    )
    titles = [action.text().replace("&", "") for action in reports.actions()]
    assert "Regenerate I/O sheets" in titles
    window.regenerate_io_sheets()
    assert window.session.project.sheet("io-bx-08nd3").component("START").library_id.endswith("pb_no")

    warnings: list[str] = []
    monkeypatch.setattr(QMessageBox, "warning", lambda *args, **kwargs: warnings.append(str(args[-1])))
    monkeypatch.setattr(PlcTagExportDialog, "exec", lambda self: self.DialogCode.Accepted)
    monkeypatch.setattr(
        PlcTagExportDialog,
        "values",
        lambda self: {
            "path": str(tmp_path / "tags.csv"),
            "target": "studio5000",
            "prefix": "",
            "name_from": "device",
            "logix_style": "alias",
            "logix_slot": 1,
        },
    )
    window.export_plc_tags()
    assert warnings
    assert "studio5000" in warnings[-1]
    window.close()


def _logix_rack() -> Session:
    """5069-IB16 slot 1, 5069-OB16 slot 2, 5069-IF8 slot 3. Forty real points."""

    session = Session.create("Rack")
    session.io_add_module("ab.5069_ib16", ref="IO1", slot=1)
    session.io_add_module("ab.5069_ob16", ref="IO2", slot=2)
    session.io_add_module("ab.5069_if8", ref="IO3", slot=3)
    session.io_assign_point("IO1", "I08", tag="XIO1_9", field="prox_3wire")
    session.io_assign_point("IO2", "Q8", tag="YOUT9", field="solenoid")
    session.io_generate_sheets()
    return session


def _render_exporter(rows: list[dict[str, str]], target: str) -> str:
    """The text each exporter writes, including targets that refuse a style mismatch."""

    from panelcad.core.plc_tags import _prepare, _render_prepared, normalize_target

    key = normalize_target(target)
    prepared = _prepare(rows, prefix="", name_from="device", logix_slot=1, analog_start=None)
    return _render_prepared(prepared, target=key, logix_style="alias")


def _numbered(prefix: str, values: list[str]) -> list[int]:
    out = []
    for value in values:
        if value.startswith(prefix) and value[len(prefix) :].isdigit():
            out.append(int(value[len(prefix) :]))
    return sorted(out)


@pytest.mark.confirm_generic_layout
def test_continuation_points_use_the_home_module(tmp_path: Path):
    import csv

    session = _logix_rack()
    rows = session.io_rows()
    assert len(rows) == 40
    assert {row["module"] for row in rows} == {"IO1", "IO2", "IO3"}
    assert {int(row["slot"]) for row in rows} == {1, 2, 3}
    by_pin = {(row["module"], row["address"]): row for row in rows}
    assert by_pin[("IO1", "I08")]["point"] == "Local:1:I.Pt08.Data"
    assert by_pin[("IO1", "I08")]["tag"] == "XIO1_9"
    assert by_pin[("IO1", "I15")]["point"] == "Local:1:I.Pt15.Data"
    assert by_pin[("IO2", "Q8")]["point"] == "Local:2:O.Pt08.Data"
    assert by_pin[("IO2", "Q8")]["tag"] == "YOUT9"
    assert by_pin[("IO2", "Q15")]["point"] == "Local:2:O.Pt15.Data"
    assert by_pin[("IO3", "AI0")]["point"] == "Local:3:I.Ch00.Data"
    assert by_pin[("IO3", "AI7")]["point"] == "Local:3:I.Ch07.Data"
    assert not any(row["module"].endswith("_2") for row in rows)

    exported = session.export_plc_tags(tmp_path / "logix.csv", target="studio5000")
    assert exported["count"] == 40
    studio = (tmp_path / "logix.csv").read_text(encoding="utf-8")
    aliases = [row for row in csv.reader(studio.splitlines()) if row and row[0] == "ALIAS"]
    assert len(aliases) == 40
    specifiers = {row[5] for row in aliases}
    assert "Local:1:I.Pt08.Data" in specifiers
    assert "Local:1:I.Pt15.Data" in specifiers
    assert "Local:2:O.Pt08.Data" in specifiers
    assert "Local:2:O.Pt15.Data" in specifiers
    assert "Local:3:I.Ch00.Data" in specifiers
    assert not any(item.startswith("Local:4:") or item.startswith("Local:5:") for item in specifiers)
    assert "Spare_I00" not in studio

    domore = _render_exporter(rows, "domore")
    elements = [row[0] for row in csv.reader(domore.splitlines()) if row]
    assert _numbered("X", elements) == list(range(16))
    assert _numbered("Y", elements) == list(range(16))
    assert _numbered("WX", elements) == list(range(8))
    assert '"X8","XIO1_9"' in domore
    assert '"Y8","YOUT9"' in domore

    keyence = _render_exporter(rows, "keyence")
    krows = [
        row
        for row in csv.reader(line for line in keyence.splitlines() if line and not line.startswith("#"))
        if row and row[0] != "name"
    ]
    assert len(krows) == 40
    kmap = {row[0]: row[1] for row in krows}
    relays = [row[1] for row in krows if row[1].startswith("R")]
    assert len(relays) == 32
    assert kmap["XIO1_9"] == "R30008"
    assert kmap["YOUT9"] == "R31008"

    for target in ("sysmac", "cxprogrammer"):
        omron = _render_exporter(rows, target)
        if target == "cxprogrammer":
            body = [line.split("\t") for line in omron.splitlines() if line]
            amap = {row[0]: row[2] for row in body}
        else:
            parsed = list(csv.reader(line for line in omron.splitlines() if line and not line.startswith("#")))
            body = parsed[1:]
            amap = {row[0]: row[2] for row in body}
        assert len(body) == 40
        assert amap["XIO1_9"] == "0.08"
        assert amap["YOUT9"] == "100.08"
        assert sum(1 for row in body if row[2] == "0.00") == 1

    only = Session.create("IB16")
    only.io_add_module("ab.5069_ib16", ref="IO1", slot=1)
    only.io_generate_sheets()
    only_rows = only.io_rows()
    assert len(only_rows) == 16
    assert {row["module"] for row in only_rows} == {"IO1"}
    assert {int(row["slot"]) for row in only_rows} == {1}
    assert any(row["address"] == "I08" and row["point"] == "Local:1:I.Pt08.Data" for row in only_rows)
    only.export_plc_tags(tmp_path / "ib16.csv", target="studio5000")
    ib_text = (tmp_path / "ib16.csv").read_text(encoding="utf-8")
    assert sum(1 for line in ib_text.splitlines() if line.startswith("ALIAS,")) == 16
    assert "Local:2:" not in ib_text

    sm = Session.create("SM1221")
    sm.io_add_module("siemens.sm1221_di16", ref="A1", address_base=8)
    sm.io_assign_point("A1", "Ib0", tag="XIO1_9", field="prox_3wire")
    sm.io_generate_sheets()
    sm_rows = sm.io_rows()
    assert len(sm_rows) == 16
    assert {row["module"] for row in sm_rows} == {"A1"}
    sm_by = {row["address"]: row for row in sm_rows}
    assert sm_by["Ia0"]["point"] == "%I8.0"
    assert sm_by["Ib0"]["point"] == "%I9.0"
    assert sm_by["Ib0"]["tag"] == "XIO1_9"
    assert sm_by["Ib7"]["point"] == "%I9.7"
    assert "Spare_Ia0" not in {row["tag"] for row in sm_rows}
    generic = sm.export_plc_tags(tmp_path / "sm.csv", target="tia_csv")
    assert generic["count"] == 16
    sm_text = (tmp_path / "sm.csv").read_text(encoding="utf-8")
    assert "%I8.0" in sm_text and "%I9.0" in sm_text and "%I9.7" in sm_text
    assert "%I2.0" not in sm_text and "Spare_Ia0" not in sm_text
    book = tmp_path / "sm.xlsx"
    workbook = sm.export_plc_tags(book, target="tia")
    assert workbook["count"] == 16
    from openpyxl import load_workbook

    table = load_workbook(book)["PLCTags"]
    addresses = [row[3] for row in table.iter_rows(min_row=2, values_only=True)]
    assert addresses[0] == "%I8.0"
    assert "%I9.0" in addresses and "%I9.7" in addresses
    assert "%I2.0" not in addresses
    assert len(addresses) == 16


@pytest.mark.confirm_generic_layout
def test_continuation_export_gui_jsonl_and_cli(tmp_path: Path, monkeypatch: pytest.MonkeyPatch):
    from PySide6.QtWidgets import QApplication, QMessageBox

    from panelcad.cli.main import app
    from panelcad.gui.main_window import MainWindow
    from panelcad.gui.plc_tag_dialog import PlcTagExportDialog

    session = _logix_rack()
    project = tmp_path / "rack.pcad.yaml"
    session.save(project)

    agent = Agent()
    assert agent.handle({"id": 1, "op": "create_project", "name": "Rack"})["ok"]
    for index, (module, ref, slot) in enumerate(
        (("ab.5069_ib16", "IO1", 1), ("ab.5069_ob16", "IO2", 2), ("ab.5069_if8", "IO3", 3)),
        start=2,
    ):
        added = agent.handle({"id": index, "op": "io_add_module", "module": module, "ref": ref, "slot": slot})
        assert added["ok"], added
    drawn = agent.handle({"id": 5, "op": "io_generate_sheets"})
    assert drawn["ok"], drawn
    listed = agent.handle({"id": 6, "op": "plc_io_list"})
    assert listed["ok"] and listed["result"]["count"] == 40
    assert {row["module"] for row in listed["result"]["rows"]} == {"IO1", "IO2", "IO3"}
    home = next(row for row in listed["result"]["rows"] if row["address"] == "I08")
    assert home["point"] == "Local:1:I.Pt08.Data"
    assert home["slot"] == "1"
    tags = tmp_path / "agent.csv"
    exported = agent.handle({"id": 7, "op": "export_plc_tags", "path": str(tags), "target": "studio5000"})
    assert exported["ok"], exported
    assert exported["result"]["count"] == 40
    assert tags.read_text(encoding="utf-8").count("ALIAS,") == 40
    assert "Local:4:" not in tags.read_text(encoding="utf-8")

    from typer.testing import CliRunner

    runner = CliRunner()
    listed_cli = runner.invoke(app, ["electrical", "plc-io", str(project)])
    assert listed_cli.exit_code == 0, listed_cli.output
    listed_modules = {line.split()[0] for line in listed_cli.output.splitlines() if line.strip()}
    assert listed_modules == {"IO1", "IO2", "IO3"}
    assert "Local:1:I.Pt08.Data" in listed_cli.output
    assert "Local:4:" not in listed_cli.output
    cli_tags = tmp_path / "cli.csv"
    cli = runner.invoke(
        app,
        ["electrical", "plc-tags", str(project), "--out", str(cli_tags), "--target", "studio5000"],
    )
    assert cli.exit_code == 0, cli.output
    assert cli_tags.read_text(encoding="utf-8").count("ALIAS,") == 40

    if QApplication.instance() is None:
        QApplication([])
    window = MainWindow(session)
    warnings: list[str] = []
    monkeypatch.setattr(QMessageBox, "warning", lambda *args, **kwargs: warnings.append(str(args[-1])))
    monkeypatch.setattr(PlcTagExportDialog, "exec", lambda self: self.DialogCode.Accepted)
    gui_tags = tmp_path / "gui.csv"
    monkeypatch.setattr(
        PlcTagExportDialog,
        "values",
        lambda self: {
            "path": str(gui_tags),
            "target": "studio5000",
            "prefix": "",
            "name_from": "device",
            "logix_style": "alias",
            "logix_slot": 1,
        },
    )
    window.export_plc_tags()
    assert warnings == []
    assert gui_tags.read_text(encoding="utf-8").count("ALIAS,") == 40
    assert "Local:1:I.Pt08.Data" in gui_tags.read_text(encoding="utf-8")
    assert "Local:4:" not in gui_tags.read_text(encoding="utf-8")
    window.close()


def _bx_started() -> Session:
    session = Session.create("BX")
    session.io_add_module("brx.bx_08nd3", ref="A1")
    session.io_assign_point("A1", "0", tag="PB_START", field="pushbutton")
    session.io_generate_sheets()
    return session


def _write_prox_csv(path: Path) -> None:
    path.write_text(
        "module,address,tag,field,description\nA1,3,PX_NEW,prox_3wire,Prox 3\n",
        encoding="utf-8",
    )


def _assert_prox_imported(session: Session) -> None:
    sheet = session.project.sheet("io-bx-08nd3")
    assert sheet.component("PX_NEW").library_id.endswith("prox_3wire")
    assert sheet.component("PB_START").library_id.endswith("pb_no")
    row = next(item for item in session.io_rows() if item["module"] == "A1" and item["address"] == "3")
    assert row["tag"] == "PX_NEW"
    assert row["field"] == "prox_3wire"


def _assert_import_undone(session: Session) -> None:
    sheet = session.project.sheet("io-bx-08nd3")
    assert "PX_NEW" not in {comp.id for comp in sheet.components}
    assert sheet.component("PB_START").library_id.endswith("pb_no")
    row = next(item for item in session.io_rows() if item["module"] == "A1" and item["address"] == "3")
    assert row["tag"] != "PX_NEW"
    assert row["field"] == "spare"
    spare = sheet.component(row["tag"])
    assert "notes.spare" in spare.library_id


def test_import_io_undo_restores_drawing_and_list(tmp_path: Path):
    session = _bx_started()
    before_rows = [(row["address"], row["tag"], row["field"]) for row in session.io_rows()]
    before_parts = [
        (sheet.id, [(comp.id, comp.library_id) for comp in sheet.components]) for sheet in session.project.sheets
    ]
    _write_prox_csv(tmp_path / "prox.csv")
    imported = session.import_io_list(tmp_path / "prox.csv")
    assert imported["errors"] == []
    _assert_prox_imported(session)
    assert session.undo()
    _assert_import_undone(session)
    assert [(row["address"], row["tag"], row["field"]) for row in session.io_rows()] == before_rows
    assert [
        (sheet.id, [(comp.id, comp.library_id) for comp in sheet.components]) for sheet in session.project.sheets
    ] == before_parts
    assert session.redo()
    _assert_prox_imported(session)


def test_import_io_undo_jsonl_and_gui(tmp_path: Path, monkeypatch: pytest.MonkeyPatch):
    from PySide6.QtWidgets import QApplication, QFileDialog

    from panelcad.gui.main_window import MainWindow

    agent = Agent()
    assert agent.handle({"id": 1, "op": "create_project", "name": "BX"})["ok"]
    added = agent.handle({"id": 2, "op": "io_add_module", "module": "brx.bx_08nd3", "ref": "A1"})
    assert added["ok"], added
    assigned = agent.handle(
        {"id": 3, "op": "io_assign_point", "module": "A1", "pin": "0", "tag": "PB_START", "field": "pushbutton"}
    )
    assert assigned["ok"], assigned
    drawn = agent.handle({"id": 4, "op": "io_generate_sheets"})
    assert drawn["ok"], drawn
    csv_path = tmp_path / "prox.csv"
    _write_prox_csv(csv_path)
    loaded = agent.handle({"id": 5, "op": "import_io", "path": str(csv_path)})
    assert loaded["ok"], loaded
    _assert_prox_imported(agent._need())
    undone = agent.handle({"id": 6, "op": "undo"})
    assert undone["ok"] and undone["result"]["undone"] is True
    _assert_import_undone(agent._need())
    redone = agent.handle({"id": 7, "op": "redo"})
    assert redone["ok"] and redone["result"]["redone"] is True
    _assert_prox_imported(agent._need())

    if QApplication.instance() is None:
        QApplication([])
    window = MainWindow(_bx_started())
    monkeypatch.setattr(QFileDialog, "getOpenFileName", lambda *args, **kwargs: (str(csv_path), "CSV (*.csv)"))
    window.import_io()
    _assert_prox_imported(window.session)
    window.undo_edit()
    _assert_import_undone(window.session)
    window.redo_edit()
    _assert_prox_imported(window.session)
    window.close()


def _wires_on(session: Session, ref: str, pin: str) -> list[str]:
    token = f"{ref}.{pin}".upper()
    numbers: list[str] = []
    for sheet in session.project.sheets:
        for wire in sheet.wires:
            ends = (str(wire.frm).upper(), str(wire.to).upper())
            if token in ends and str(wire.number or "").strip():
                numbers.append(str(wire.number))
    return numbers


def _continuation_ids(session: Session, home: str) -> list[str]:
    found = []
    for sheet in session.project.sheets:
        for comp in sheet.components:
            if str(comp.properties.get("io_continuation") or "") == home:
                found.append(comp.id)
    return found


def _assert_continuation_landed(session: Session, module: str, pin: str, tag: str, point: str) -> str:
    """The home row keeps the real address and the wire drawn on the continuation card."""

    rows = {(row["module"], row["address"]): row for row in session.io_rows()}
    row = rows[(module, pin)]
    assert row["point"] == point
    assert row["tag"] == tag
    assert row["spare"] == "no"
    assert row["device"] == tag
    assert row["field"] != "spare"
    home_wires = _wires_on(session, module, pin)
    assert home_wires == []
    cont_wires: list[str] = []
    for cont_id in _continuation_ids(session, module):
        cont_wires.extend(_wires_on(session, cont_id, pin))
    assert cont_wires
    assert row["wire"] == cont_wires[0]
    assert row["wire"]
    return row["wire"]


@pytest.mark.confirm_generic_layout
def test_continuation_wires_stay_with_the_home_row(tmp_path: Path):
    import csv

    logix = Session.create("Logix")
    logix.io_add_module("ab.5069_ib16", ref="IO1", slot=1)
    logix.io_add_module("ab.5069_ob16", ref="IO2", slot=2)
    logix.io_assign_point("IO1", "I08", tag="XIO1_9", description="Prox 9", field="prox_3wire")
    logix.io_assign_point("IO1", "I15", tag="IB_15", description="Prox 15", field="prox_3wire")
    logix.io_assign_point("IO2", "Q8", tag="YOUT9", description="Sol 9", field="solenoid")
    logix.io_assign_point("IO2", "Q15", tag="PL_RUN", description="Run", field="pilot_light")
    logix.io_generate_sheets()
    rows = logix.io_rows()
    assert len(rows) == 32
    assert {row["module"] for row in rows} == {"IO1", "IO2"}
    wire_i08 = _assert_continuation_landed(logix, "IO1", "I08", "XIO1_9", "Local:1:I.Pt08.Data")
    _assert_continuation_landed(logix, "IO1", "I15", "IB_15", "Local:1:I.Pt15.Data")
    wire_q8 = _assert_continuation_landed(logix, "IO2", "Q8", "YOUT9", "Local:2:O.Pt08.Data")
    _assert_continuation_landed(logix, "IO2", "Q15", "PL_RUN", "Local:2:O.Pt15.Data")
    for pin in (f"I{index:02d}" for index in range(8, 16)):
        assert rows_pin(rows, "IO1", pin)["wire"]
    for pin in (f"Q{index}" for index in range(8, 16)):
        assert rows_pin(rows, "IO2", pin)["wire"]

    logix.export_plc_tags(tmp_path / "logix.csv", target="studio5000")
    studio = (tmp_path / "logix.csv").read_text(encoding="utf-8")
    aliases = [row for row in csv.reader(studio.splitlines()) if row and row[0] == "ALIAS"]
    assert len(aliases) == 32
    by_name = {row[2]: row[5] for row in aliases}
    assert by_name["XIO1_9"] == "Local:1:I.Pt08.Data"
    assert by_name["PL_RUN"] == "Local:2:O.Pt15.Data"
    assert not any(item.startswith("Local:3:") for item in by_name.values())

    domore_text = _render_exporter(rows, "domore")
    domore_rows = {row[1]: row for row in csv.reader(domore_text.splitlines()) if row}
    assert domore_rows["XIO1_9"][0] == "X8"
    assert domore_rows["XIO1_9"][2] == wire_i08
    assert domore_rows["YOUT9"][0] == "Y8"
    assert domore_rows["YOUT9"][2] == wire_q8

    keyence = _render_exporter(rows, "keyence")
    krows = [
        row
        for row in csv.reader(line for line in keyence.splitlines() if line and not line.startswith("#"))
        if row and row[0] != "name"
    ]
    kmap = {row[0]: row[1] for row in krows}
    assert len(krows) == 32
    assert kmap["XIO1_9"] == "R30008"
    assert kmap["YOUT9"] == "R31008"
    assert kmap["PL_RUN"] == "R31015"

    for target in ("sysmac", "cxprogrammer"):
        omron = _render_exporter(rows, target)
        if target == "cxprogrammer":
            body = [line.split("\t") for line in omron.splitlines() if line]
        else:
            parsed = list(csv.reader(line for line in omron.splitlines() if line and not line.startswith("#")))
            body = parsed[1:]
        amap = {row[0]: row[2] for row in body}
        assert len(body) == 32
        assert amap["XIO1_9"] == "0.08"
        assert amap["YOUT9"] == "100.08"
        assert amap["PL_RUN"] == "100.15"

    tia = Session.create("TIA")
    tia.io_add_module("siemens.sm1221_di16", ref="A1", address_base=8)
    tia.io_assign_point("A1", "Ib7", tag="LS_DOOR", description="Door", field="limit_switch")
    tia.io_generate_sheets()
    assert len(tia.io_rows()) == 16
    _assert_continuation_landed(tia, "A1", "Ib7", "LS_DOOR", "%I9.7")
    tia.export_plc_tags(tmp_path / "tia.csv", target="tia_csv")
    tia_text = (tmp_path / "tia.csv").read_text(encoding="utf-8")
    assert "%I9.7" in tia_text
    assert "%I2.0" not in tia_text
    assert "Spare_Ia0" not in tia_text

    brx = Session.create("BRX")
    brx.io_add_module("brx.bx_08nd3", ref="X1")
    brx.io_add_module("brx.bx_12nd3", ref="X2", polarity="source")
    brx.io_assign_point("X2", "8", tag="PX_SRC", description="Source prox", field="prox_3wire")
    brx.io_assign_point("X2", "11", tag="LS_SRC", description="Source limit", field="limit_switch")
    brx.io_generate_sheets()
    wire_x16 = _assert_continuation_landed(brx, "X2", "8", "PX_SRC", "X16")
    _assert_continuation_landed(brx, "X2", "11", "LS_SRC", "X19")
    brx.export_plc_tags(tmp_path / "brx.csv", target="domore")
    exported = {row[1]: row for row in csv.reader((tmp_path / "brx.csv").read_text(encoding="utf-8").splitlines()) if row}
    assert exported["PX_SRC"][0] == "X16"
    assert exported["PX_SRC"][2] == wire_x16
    assert exported["LS_SRC"][0] == "X19"

    micro = Session.create("M8")
    micro.io_add_module("ab.2085_iq16", ref="M1", slot=2)
    micro.io_assign_point("M1", "PT09", tag="PX_JOG", description="Jog prox", field="prox_3wire")
    micro.io_generate_sheets()
    _assert_continuation_landed(micro, "M1", "PT09", "PX_JOG", "_IO_X2_DI_09")


def rows_pin(rows: list[dict[str, str]], module: str, pin: str) -> dict[str, str]:
    return next(row for row in rows if row["module"] == module and row["address"] == pin)


@pytest.mark.confirm_generic_layout
def test_continuation_wire_gui_jsonl_and_cli(tmp_path: Path, monkeypatch: pytest.MonkeyPatch):
    import csv

    from PySide6.QtWidgets import QApplication, QFileDialog

    from panelcad.cli.main import app
    from panelcad.gui.main_window import MainWindow

    session = Session.create("Logix")
    session.io_add_module("ab.5069_ib16", ref="IO1", slot=1)
    session.io_add_module("ab.5069_ob16", ref="IO2", slot=2)
    session.io_assign_point("IO1", "I08", tag="XIO1_9", description="Prox 9", field="prox_3wire")
    session.io_assign_point("IO2", "Q8", tag="YOUT9", description="Sol 9", field="solenoid")
    session.io_generate_sheets()
    wire = _assert_continuation_landed(session, "IO1", "I08", "XIO1_9", "Local:1:I.Pt08.Data")
    project = tmp_path / "rack.pcad.yaml"
    session.save(project)

    agent = Agent()
    opened = agent.handle({"id": 1, "op": "open", "path": str(project)})
    assert opened["ok"], opened
    listed = agent.handle({"id": 2, "op": "plc_io_list"})
    assert listed["ok"], listed
    hit = next(row for row in listed["result"]["rows"] if row["tag"] == "XIO1_9")
    assert hit["spare"] == "no"
    assert hit["wire"] == wire
    assert hit["device"] == "XIO1_9"
    assert hit["module"] == "IO1"
    assert hit["point"] == "Local:1:I.Pt08.Data"
    assert listed["result"]["count"] == 32

    from typer.testing import CliRunner

    runner = CliRunner()
    text = runner.invoke(app, ["electrical", "plc-io", str(project)])
    assert text.exit_code == 0, text.output
    line = next(item for item in text.output.splitlines() if "XIO1_9" in item)
    assert " no " in f" {line} "
    assert wire in line
    assert "Local:1:I.Pt08.Data" in line
    out = tmp_path / "io.csv"
    saved = runner.invoke(app, ["electrical", "plc-io", str(project), "--out", str(out)])
    assert saved.exit_code == 0, saved.output
    csv_rows = list(csv.DictReader(out.read_text(encoding="utf-8").splitlines()))
    csv_hit = next(row for row in csv_rows if row["tag"] == "XIO1_9")
    assert csv_hit["spare"] == "no"
    assert csv_hit["wire"] == wire
    assert csv_hit["device"] == "XIO1_9"

    if QApplication.instance() is None:
        QApplication([])
    window = MainWindow(session)
    gui_csv = tmp_path / "gui-io.csv"
    monkeypatch.setattr(QFileDialog, "getSaveFileName", lambda *args, **kwargs: (str(gui_csv), "CSV (*.csv)"))
    reports = next(
        action.menu()
        for action in window.menuBar().actions()
        if action.text().replace("&", "") == "Reports"
    )
    export = next(action for action in reports.actions() if "I/O List" in action.text().replace("&", ""))
    export.trigger()
    gui_rows = list(csv.DictReader(gui_csv.read_text(encoding="utf-8").splitlines()))
    gui_hit = next(row for row in gui_rows if row["tag"] == "XIO1_9")
    assert gui_hit["spare"] == "no"
    assert gui_hit["wire"] == wire
    assert gui_hit["device"] == "XIO1_9"
    assert gui_hit["point"] == "Local:1:I.Pt08.Data"
    window.close()
